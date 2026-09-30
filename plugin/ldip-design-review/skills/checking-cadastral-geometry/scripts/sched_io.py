"""Shared reader for Parsons masterplan schedules (xlsx or csv) — skill 9 checking-cadastral-geometry.

Reads a sheet into a grid of raw values with Excel-style cell references, parses the number
formats seen in the LDIP schedules (1,263.04 · " -   " · (88.64) · 0.34%), and locates
header columns by label (SPC-0004 §5.2: header located by its labels, never by position).
No network, no shell calls.
"""
__version__ = "0.1.0"

import csv
import hashlib
import re
from pathlib import Path

EXIT_OK, EXIT_INPUT, EXIT_DEPENDENCY, EXIT_SCHEMA, EXIT_REFUSED = 0, 2, 3, 4, 5


class SchedError(Exception):
    def __init__(self, code, detail):
        super().__init__(detail)
        self.code, self.detail = code, detail


def col_letter(idx0):
    s, n = "", idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def cell_ref(row1, col0):
    return f"{col_letter(col0)}{row1}"


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


MAX_PATH = 250   # SPC-0002 R1; SPC-0006 §2.3.4


def check_path(path):
    """Exit 2 on a path longer than 250 characters (SPC-0006 §2.3.4)."""
    if len(str(Path(path).resolve())) > MAX_PATH:
        raise SchedError(EXIT_INPUT, f"path longer than {MAX_PATH} characters: {path}")
    return Path(path)


def load_grid(path, sheet=None):
    """Return (grid, source) where grid[r][c] holds raw values (row 0 = Excel row 1)."""
    p = check_path(path)
    if not p.is_file():
        raise SchedError(EXIT_INPUT, f"file not found: {p}")
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError as exc:
            raise SchedError(EXIT_DEPENDENCY, "openpyxl is not installed") from exc
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            raise SchedError(EXIT_INPUT, f"sheet '{sheet}' not in workbook (sheets: {wb.sheetnames})")
        grid = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
        wb.close()
        src_type = "xlsx_cell"
    elif ext == ".csv":
        with open(p, newline="", encoding="utf-8-sig", errors="replace") as fh:
            grid = [list(r) for r in csv.reader(fh)]
        src_type = "csv_cell"
    else:
        raise SchedError(EXIT_INPUT, f"unsupported file type {ext} (xlsx or csv)")
    width = max((len(r) for r in grid), default=0)
    grid = [r + [None] * (width - len(r)) for r in grid]
    return grid, {"path": str(p), "sha256": sha256_file(p), "sheet": sheet if src_type == "xlsx_cell" else None,
                  "source_type": src_type, "file_type": ext.lstrip(".")}


_NUM = re.compile(r"^\(?-?[0-9]*\.?[0-9]+\)?%?$")


def to_num(v):
    """Parse a schedule value to float. '-' / '' → 0.0 in numeric columns is left to the caller: this
    returns None for blank and the string '-'. Returns None when not a number."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace(" ", "")
    if s in ("", "-", "–"):
        return None
    if not _NUM.match(s):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    pct = s.endswith("%")
    s = s.rstrip("%")
    try:
        x = float(s)
    except ValueError:
        return None
    if pct:
        x /= 100.0
    return -x if neg else x


def is_dash(v):
    return isinstance(v, str) and v.strip() in ("-", "–")


def norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def find_label(grid, label, max_rows=40):
    """First (row0, col0) whose normalised text equals label."""
    target = norm(label)
    for r in range(min(max_rows, len(grid))):
        for c, v in enumerate(grid[r]):
            if norm(v) == target:
                return r, c
    return None


def header_map(grid, top_row0, rows=2):
    """Map normalised header labels to column index over `rows` header rows. A sub-header is keyed
    'parent/child' using the nearest non-empty parent to its left (merged header cells)."""
    out = {}
    top = grid[top_row0]
    parents, last = [], None
    for c in range(len(top)):
        if top[c] not in (None, ""):
            last = norm(top[c])
        parents.append(last)
    for c, v in enumerate(top):
        if v not in (None, ""):
            out.setdefault(norm(v), c)
    for k in range(1, rows):
        if top_row0 + k >= len(grid):
            break
        for c, v in enumerate(grid[top_row0 + k]):
            if v not in (None, ""):
                out.setdefault(norm(v), c)
                if parents[c]:
                    out.setdefault(f"{parents[c]}/{norm(v)}", c)
    return out
