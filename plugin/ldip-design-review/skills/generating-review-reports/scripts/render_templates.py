#!/usr/bin/env python3
"""Render skill 22 drafts into the CRS and RFI forms (skill 24, SPC-0006 §4.8.3 / §4.10.3; P-S22-6).

Usage: python render_templates.py --drafts <drafts.json> --out-dir <folder>
           [--crs-template <CRS form .xlsx>] [--rfi-template <RFI form .docx>]
Outputs, all marked DRAFT with the plugin version:
  CRS-<SITE>-<sub>-DRAFT.xlsx   one per CRS group of the drafts; the CRS form, plus a "Draft trace" sheet
                                (finding id, rule, status, severity, route basis, EOT flag) for AICC use
  RFI-<SITE>-<nn>-DRAFT.docx    one per RFI draft, in the RFI form
The forms are the owner's files, kept outside Git (they hold project data). Fields are found by their
printed labels, not by fixed cells; every value in the form (sample comments, names, dates, ticks) is
cleared before the draft is written, and the source files are never changed. Without a template a neutral
form with the same labels and columns is generated. Comment-level status codes, review codes, reviewer
initials, signatures and dates are left empty for the reviewers. Nothing is issued.
Exit: 0 written · 2 input error (template unreadable, labels not found) · 3 openpyxl / python-docx missing.
No network, no shell calls.
"""
__version__ = "0.1.0"

import argparse
import copy
import math
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_common import (DRAFT, EXIT_DEP, EXIT_INPUT, EXIT_OK, atomic_save, emit, fail, path_ok,  # noqa: E402
                           plugin_version, read_json, safe_name)

COLS = ["#", "Page/Para/Clause/ Doc. No. Etc.", "Review Comment", "Reviewer\nInitial", "Review No. ",
        "Contractor's Response", "Responder\nInitial", "Status Code", "Comment Status (Closed/Open)"]
# CRS header labels (column B and column F of the form) -> drafts header key ("" = cleared)
CRS_LEFT = {"crs ref": "crs_ref", "program": "program", "project": "project", "section": "section",
            "sub-section": "sub_section", "pmcm": "pmcm", "transmittal no.": "", "subject": "subject",
            "document no": "document_no", "reviewer/s organization": "reviewer_org", "reviewer/s name": "",
            "contractor": "contractor"}
CRS_RIGHT = {"this submission is returned (review code)": "Select Status Code",
             "date comments issued to contractor": "", "date of receiving contractor's response": ""}


class TemplateError(Exception):
    pass


def norm(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().rstrip(":").strip().lower()


# ---------------------------------------------------------------- CRS (.xlsx)

def neutral_crs():
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    thin = Side(style="thin")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CRS"
    ws["A5"] = "Comments Review Sheet (CRS)"
    ws["A5"].font = Font(bold=True, size=14)
    left = ["CRS Ref:", "Program", "Project:", "Section", "Sub-Section", "PMCM", "Transmittal No.:", "Subject:",
            "Document No:", "Reviewer/s Organization:", "Reviewer/s Name:", "Contractor:"]
    right = ["This submission is returned (Review Code):", "Date Comments Issued to Contractor: ",
             "Date of receiving Contractor's Response: "]
    for i, lab in enumerate(left):
        ws.cell(row=6 + i, column=2, value=lab).font = Font(bold=True)
    for i, lab in enumerate(right):
        ws.cell(row=7 + i, column=6, value=lab).font = Font(bold=True)
    for j, h in enumerate(COLS, 1):
        c = ws.cell(row=20, column=j, value=h)
        c.font = Font(bold=True)
        c.border = Border(thin, thin, thin, thin)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r in range(21, 37):
        for j in range(1, 10):
            ws.cell(row=r, column=j).border = Border(thin, thin, thin, thin)
            ws.cell(row=r, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    for i in range(3):
        ws.cell(row=38 + i, column=1, value=f"Reviewer {i + 1}: ")
        ws.cell(row=38 + i, column=5, value="Sign:  ")
        ws.cell(row=38 + i, column=7, value="Date: ")
    for col, w in zip("ABCDEFGHI", [6.5, 43, 72.5, 12.5, 10, 70.5, 26.5, 14.5, 19.5]):
        ws.column_dimensions[col].width = w
    return wb


def find_form_sheet(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(max_row=30):
            for c in row:
                if norm(c.value) == "#" and c.column == 1:
                    return ws, c.row
    raise TemplateError("CRS form: no comment table header ('#' in column A) found")


def shift_rows(ws, idx, k):
    """Insert k rows before row idx, moving merged cells, row heights and the print area (openpyxl does not)."""
    merged = [m for m in ws.merged_cells.ranges if m.min_row >= idx]
    for m in merged:
        ws.unmerge_cells(str(m))
    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= idx and d.height}
    ws.insert_rows(idx, k)
    for r in sorted(heights, reverse=True):
        ws.row_dimensions[r + k].height = heights[r]
        ws.row_dimensions[r].height = None
    for m in merged:
        m.shift(0, k)
        ws.merge_cells(str(m))
    pa = ws.print_area
    if pa:
        m = re.search(r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", pa if isinstance(pa, str) else str(pa))
        if m and int(m.group(4)) >= idx:
            ws.print_area = f"{m.group(1)}{m.group(2)}:{m.group(3)}{int(m.group(4)) + k}"


def row_height(text, width_chars):
    lines = sum(max(1, math.ceil(len(part) / max(width_chars, 10))) for part in str(text).split("\n"))
    return max(30, min(409, 16 * lines + 8))


def fill_crs(wb, group, marking, version):
    from openpyxl.styles import Font
    ws, head = find_form_sheet(wb)
    footer = next((r for r in range(head + 1, ws.max_row + 1)
                   if norm(ws.cell(r, 1).value).startswith("reviewer")), None)
    if footer is None:
        raise TemplateError("CRS form: no reviewer block ('Reviewer 1:' in column A) below the table")
    h = group["header"]
    found = set()
    for r in range(1, head):
        for col, table in ((2, CRS_LEFT), (6, CRS_RIGHT)):
            lab = norm(ws.cell(r, col).value)
            if lab in table:
                key = table[lab]
                val = (h.get(key, "") if key else "") if col == 2 else key
                target = ws.cell(r, col + 1)
                if type(target).__name__ == "MergedCell":
                    continue
                target.value = val or None
                found.add(lab)
    missing = {"subject", "document no", "crs ref"} - found
    if missing:
        raise TemplateError(f"CRS form: header labels not found: {sorted(missing)}")
    # comment rows: from the first numbered row (or the row under the header) to the last bordered row
    first = next((r for r in range(head + 1, footer) if isinstance(ws.cell(r, 1).value, (int, float))), head + 1)
    last = max((r for r in range(first, footer) if ws.cell(r, 3).border.left.style), default=first)
    need = len(group["rows"])
    cap = last - first + 1
    if need > cap:
        style_row = last
        shift_rows(ws, last + 1, need - cap)
        for r in range(last + 1, last + 1 + need - cap):
            for j in range(1, 10):
                ws.cell(r, j)._style = copy.copy(ws.cell(style_row, j)._style)
        for dv in ws.data_validations.dataValidation:
            for rng in list(dv.sqref.ranges):
                if rng.min_row <= last <= rng.max_row:
                    rng.expand(down=need - cap)
        last += need - cap
        footer += need - cap
    for r in range(first, last + 1):
        for j in range(1, 10):
            c = ws.cell(r, j)
            if type(c).__name__ != "MergedCell":
                c.value = None
    width = int((ws.column_dimensions["C"].width or 72) * 0.85)
    styles = [copy.copy(ws.cell(first, j)._style) for j in range(1, 10)]   # one look for every comment row
    for i, row in enumerate(group["rows"]):
        r = first + i
        for j in range(1, 10):
            if type(ws.cell(r, j)).__name__ != "MergedCell":
                ws.cell(r, j)._style = copy.copy(styles[j - 1])
        for j, v in ((1, row["no"]), (2, row["location"]), (3, row["comment"]), (5, row["review_no"]),
                     (9, row["comment_status"])):
            ws.cell(r, j).value = v
        ws.row_dimensions[r].height = row_height(row["comment"], width)
    for r in range(first + need, last + 1):
        ws.row_dimensions[r].height = 18
    for r in range(footer, ws.max_row + 1):
        if norm(ws.cell(r, 1).value).startswith("reviewer"):
            for j in range(2, 10):
                c = ws.cell(r, j)
                if type(c).__name__ == "MergedCell":
                    continue
                v = norm(c.value)
                c.value = "Sign:  " if v.startswith("sign") else ("Date: " if v.startswith("date") else None)
    # DRAFT marking: the empty row above the table, the page header and footer
    mark_row = head - 1 if not any(ws.cell(head - 1, j).value for j in range(1, 10)) else head
    if mark_row == head - 1:
        ws.cell(mark_row, 1).value = marking
        ws.cell(mark_row, 1).font = Font(bold=True, color="C00000", size=12)
    ws.oddHeader.center.text = DRAFT
    ws.oddFooter.center.text = f"ldip-design-review {plugin_version()} — render_templates {version}"
    ws.sheet_properties.tabColor = "C00000"
    tr = wb.create_sheet("Draft trace")
    tr["A1"] = marking
    tr["A1"].font = Font(bold=True, color="C00000")
    tr["A2"] = "For AICC use only — remove this sheet before issue. Run(s): " + ", ".join(group["run_ids"])
    tr.append(["#", "Finding id", "Rule", "Status", "Severity", "Route basis", "EOT assessment"])
    for c in tr[3]:
        c.font = Font(bold=True)
    for row in group["rows"]:
        tr.append([row["no"], row["finding_id"], row["rule_id"], row["status"], row["severity"],
                   row["route_basis"], "yes" if row["eot_assessment"] else ""])
    for col, w in zip("ABCDEFG", [5, 52, 10, 20, 12, 46, 14]):
        tr.column_dimensions[col].width = w
    return first, last


def render_crs(group, template, out, marking):
    import openpyxl
    if template:
        try:
            wb = openpyxl.load_workbook(template)
        except Exception as exc:  # openpyxl raises several types for a damaged file
            raise TemplateError(f"CRS template unreadable: {exc}") from exc
    else:
        wb = neutral_crs()
    fill_crs(wb, group, marking, __version__)
    atomic_save(wb, out)


# ---------------------------------------------------------------- RFI (.docx)

def neutral_rfi():
    import docx
    d = docx.Document()
    t = d.add_table(rows=5, cols=4)
    t.style = "Table Grid"
    t.cell(0, 0).merge(t.cell(0, 3)).text = "REQUEST FOR INFORMATION (RFI)"
    for r, (a, b) in enumerate([("RFI Reference Number:", "Rev. No."), ("Discipline:", "Issue Date:"),
                                ("Enclosure/s:", "Reply Required Date:"), ("Priority:", "Reply Required Date:")], 1):
        t.cell(r, 0).text, t.cell(r, 2).text = a, b
    for lab, rows in (("SUBJECT:", None), ("RFI REFERENCES:", ["Drawing No(s).:", "Document:", "Specification:",
                                                                 "Location:"]),
                      ("REASON FOR RFI:", "checks"), ("INFORMATION REQUESTED:", None), ("SUBMITTED BY:", "contractor")):
        d.add_paragraph(lab)
        if rows is None:
            d.add_table(rows=1, cols=1).style = "Table Grid"
        elif rows == "checks":
            tb = d.add_table(rows=2, cols=3)
            tb.style = "Table Grid"
            tb.cell(0, 0).text, tb.cell(0, 1).text, tb.cell(0, 2).text = (
                "☐ Insufficient Information", "☐ Engineering Conflict", "☐ Alternate Proposal")
            tb.cell(1, 0).merge(tb.cell(1, 2)).text = "☐ Other, provide explanation:"
        elif rows == "contractor":
            tb = d.add_table(rows=2, cols=2)
            tb.style = "Table Grid"
            tb.cell(0, 0).text, tb.cell(0, 1).text = "Contractor:", "JV LADUN & AL-AYUNI"
            tb.cell(1, 0).text, tb.cell(1, 1).text = "Signature", "Date:"
        else:
            tb = d.add_table(rows=len(rows), cols=2)
            tb.style = "Table Grid"
            for i, lab2 in enumerate(rows):
                tb.cell(i, 0).text = lab2
    d.add_paragraph("RESPONSE:")
    tb = d.add_table(rows=2, cols=2)
    tb.style = "Table Grid"
    tb.cell(0, 0).text, tb.cell(0, 1).text = "Designer Response:", "PMCM Response:"
    tb = d.add_table(rows=3, cols=2)
    tb.style = "Table Grid"
    tb.cell(0, 0).merge(tb.cell(0, 1)).text = "ATTACHED DOCUMENTS:"
    return d


def set_text(cell, text):
    """Replace a table cell's text, keeping the first run's formatting."""
    paras = cell.paragraphs
    p = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def uniq_cells(row):
    out, prev = [], None
    for c in row.cells:
        if c._tc is not prev:
            out.append(c)
        prev = c._tc
    return out


def body_blocks(d):
    """Paragraphs and tables of the body in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            yield Paragraph(el, d)
        elif tag == "tbl":
            yield Table(el, d)


def fill_rfi(d, item, marking):
    from docx.shared import Pt, RGBColor
    from docx.table import Table
    values = {"rfi reference number": item["rfi_ref"], "rev. no.": item["rev_no"], "discipline": item["discipline"],
              "issue date": "", "enclosure/s": "Yes" if item["attachments"] else "No", "priority": "",
              "reply required date": "", "drawing no(s).": item["drawing_nos"], "document": item["document"],
              "specification": item["specification"], "location": item["location"]}
    found, blocks = set(), list(body_blocks(d))
    for b in blocks:
        if not isinstance(b, Table):
            continue
        for row in b.rows:
            cells = uniq_cells(row)
            for i, c in enumerate(cells):
                lab = norm(c.text)
                if lab in values and i + 1 < len(cells):
                    set_text(cells[i + 1], values[lab])
                    found.add(lab)
                elif "☒" in c.text:
                    set_text(c, c.text.replace("☒", "☐"))
                elif lab.startswith("date") and len(lab) > 4:
                    set_text(c, "Date:")
    after = {"subject": item["subject"], "information requested": item["information_requested"]}
    for i, b in enumerate(blocks):
        if isinstance(b, Table) or norm(b.text) not in ("subject", "information requested", "response",
                                                       "attached documents"):
            continue
        key = norm(b.text)
        nxt = next((x for x in blocks[i + 1:] if isinstance(x, Table)), None)
        if nxt is None:
            continue
        if key in after:
            set_text(nxt.cell(0, 0), after[key])
            found.add(key)
    for b in blocks:  # responses and attachments: clear, then list evidence files
        if not isinstance(b, Table):
            continue
        texts = [norm(c.text) for c in uniq_cells(b.rows[0])]
        if "designer response" in texts:
            for row in b.rows[1:]:
                for c in uniq_cells(row):
                    if not any(norm(c.text).startswith(x) for x in ("designer stamp", "pmcm stamp", "note")):
                        set_text(c, "")
        if texts and texts[0] == "attached documents":
            cells = [c for row in b.rows[1:] for c in uniq_cells(row)]
            for c in cells:
                set_text(c, "")
            for c, path in zip(cells, item["attachments"]):
                set_text(c, path)
    missing = {"rfi reference number", "subject", "information requested", "drawing no(s).", "location"} - found
    if missing:
        raise TemplateError(f"RFI form: labels not found: {sorted(missing)}")
    for s in d.sections:
        hp = s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph()
        new = hp.insert_paragraph_before("")
        run = new.add_run(marking)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(9), RGBColor(0xC0, 0, 0)
    first = blocks[0]
    para = (first.insert_paragraph_before("") if not isinstance(first, Table) else None)
    if para is None:
        from docx.oxml import OxmlElement
        p_el = OxmlElement("w:p")
        first._element.addprevious(p_el)
        from docx.text.paragraph import Paragraph
        para = Paragraph(p_el, d)
    run = para.add_run(f"{DRAFT} — finding {item['finding_id']}")
    run.bold, run.font.color.rgb = True, RGBColor(0xC0, 0, 0)


def render_rfi(item, template, out, marking):
    import docx
    if template:
        try:
            d = docx.Document(template)
        except Exception as exc:  # python-docx raises several types for a damaged file
            raise TemplateError(f"RFI template unreadable: {exc}") from exc
    else:
        d = neutral_rfi()
    fill_rfi(d, item, marking)
    atomic_save(d, out)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--drafts", required=True)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--crs-template")
    ap.add_argument("--rfi-template")
    a = ap.parse_args(argv)
    for p in [a.drafts, a.out_dir, *[x for x in (a.crs_template, a.rfi_template) if x]]:
        if not path_ok(p):
            return fail(EXIT_INPUT, f"path longer than 250 characters: {p}")
    for t, ext in ((a.crs_template, ".xlsx"), (a.rfi_template, ".docx")):
        if t and (not Path(t).is_file() or Path(t).suffix.lower() != ext):
            return fail(EXIT_INPUT, f"template not found or not {ext}: {t}")
    try:
        import docx  # noqa: F401
        import openpyxl  # noqa: F401
    except ImportError as exc:
        return fail(EXIT_DEP, f"missing package: {exc.name} (openpyxl and python-docx are needed)")
    try:
        drafts = read_json(a.drafts)
    except (OSError, ValueError) as exc:
        return fail(EXIT_INPUT, f"cannot read {a.drafts}: {exc}")
    if drafts.get("drafts_version") != "1.0" or not str(drafts.get("marking", "")).startswith(DRAFT):
        return fail(EXIT_INPUT, "not a drafts file from draft_crs.py (version or DRAFT marking missing)")
    out_dir = Path(a.out_dir)
    if a.crs_template and out_dir.resolve() == Path(a.crs_template).resolve().parent:
        return fail(EXIT_INPUT, "write the drafts to another folder than the templates' folder")
    written = []
    try:
        for g in drafts["crs"]:
            out = out_dir / f"CRS-{g['site']}-{safe_name(g['sub_id'])}-DRAFT.xlsx"
            render_crs(g, a.crs_template, out, drafts["marking"])
            written.append(str(out))
        for item in drafts["rfi"]:
            out = out_dir / f"RFI-{item['site']}-{item['no']:02d}-DRAFT.docx"
            render_rfi(item, a.rfi_template, out, drafts["marking"])
            written.append(str(out))
    except TemplateError as exc:
        return fail(EXIT_INPUT, str(exc), written=written)
    emit({"status": "ok", "crs_files": len(drafts["crs"]), "rfi_files": len(drafts["rfi"]),
          "crs_template": "owner form" if a.crs_template else "neutral",
          "rfi_template": "owner form" if a.rfi_template else "neutral", "written": written})
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
