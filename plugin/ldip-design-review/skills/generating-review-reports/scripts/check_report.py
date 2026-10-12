#!/usr/bin/env python3
"""Check that every count in the DRAFT report equals the count recomputed from the findings files
(skill 24, SPC-0006 §4.10.3 / T3).

Usage: python check_report.py --view <findings_view.json> --md <report.md> --xlsx <report.xlsx>
The findings files named in the view are read again from disk (their SHA-256 must equal the view's) and
counted again here, without using the view's rows: total, status x severity, routes and findings per
rule. Those numbers are compared with the Markdown tables and with the Summary and "By rule" sheets of
the workbook. Any difference is listed; nothing is corrected.
Exit: 0 pass · 2 input error · 3 openpyxl missing · 4 a count differs or a file changed.
No network, no shell calls.
"""
__version__ = "0.1.0"

import argparse
import collections
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_common import (EXIT_DEP, EXIT_INPUT, EXIT_OK, EXIT_SCHEMA, SEVERITIES, STATUSES, UNVERIFIED,  # noqa: E402
                           emit, fail, path_ok, provisional_route, read_json, sha256_file, status_of)
from load_findings import legacy_lists  # noqa: E402

MATRIX_HEAD = "Status \\ Severity"


def recount(view):
    problems, rows = [], []
    for i in view["inputs"]:
        p = i["path"]
        if not Path(p).is_file():
            problems.append(f"input file missing: {p}")
            continue
        if sha256_file(p) != i["sha256"]:
            problems.append(f"input file changed since the view was built: {p}")
        data = read_json(p)
        rows += data.get("findings") if isinstance(data, dict) else data
    legacy = legacy_lists(view["rulebook"]["path"])[0] if view.get("rulebook") else []
    m = collections.Counter((status_of(f) or UNVERIFIED, f["severity"]) for f in rows)
    matrix = {s: {v: m[(s, v)] for v in SEVERITIES} for s in [*STATUSES, UNVERIFIED]}
    routes = collections.Counter(provisional_route(f, legacy)[0] for f in rows)
    rules = collections.Counter((f.get("reference") or {}).get("rule_id") or "(no rule)" for f in rows)
    return {"total": len(rows), "matrix": matrix, "routes": dict(routes), "rules": dict(rules)}, problems


def md_tables(text):
    """{section heading: [rows as lists]} for every Markdown table under a '## ' heading."""
    out, head = {}, None
    for line in text.splitlines():
        if line.startswith("## "):
            head = line[3:].strip()
        elif line.startswith("|") and head and not line.startswith("|---"):
            cells = [c.strip() for c in line.strip().strip("|").split(" | ")]
            out.setdefault(head, []).append(cells)
    return out


def compare(label, expected, got, problems):
    if str(expected) != str(got):
        problems.append(f"{label}: report shows {got}, recount gives {expected}")


def check_matrix(src, rows, c, problems):
    head = rows[0]
    if head[:len(SEVERITIES) + 1] != [MATRIX_HEAD, *SEVERITIES]:
        problems.append(f"{src}: status x severity table header not found")
        return
    seen = set()
    for r in rows[1:]:
        s = r[0]
        if s == "Total":
            for j, v in enumerate(SEVERITIES, 1):
                compare(f"{src} total {v}", sum(c["matrix"][x][v] for x in c["matrix"]), r[j], problems)
            compare(f"{src} total", c["total"], r[len(SEVERITIES) + 1], problems)
            continue
        if s not in c["matrix"]:
            problems.append(f"{src}: unknown status row '{s}'")
            continue
        seen.add(s)
        for j, v in enumerate(SEVERITIES, 1):
            compare(f"{src} {s} / {v}", c["matrix"][s][v], r[j], problems)
        compare(f"{src} {s} total", sum(c["matrix"][s].values()), r[len(SEVERITIES) + 1], problems)
    for s in set(c["matrix"]) - seen:
        problems.append(f"{src}: status row '{s}' missing")


def check_routes(src, rows, c, problems):
    shown = {r[0]: r[1] for r in rows[1:]}
    for k in set(c["routes"]) | set(shown):
        compare(f"{src} route {k}", c["routes"].get(k, 0), shown.get(k, "missing"), problems)


def sheet_block(ws, first_header):
    """Rows (as strings) of the table whose header row starts with first_header, up to the first empty row."""
    rows, on = [], False
    for r in ws.iter_rows(values_only=True):
        vals = ["" if v is None else str(v) for v in r]
        if not on and vals and vals[0] == first_header:
            on = True
        if on:
            if not any(vals):
                break
            rows.append([v for v in vals])
    return rows


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--view", required=True)
    ap.add_argument("--md", required=True)
    ap.add_argument("--xlsx", required=True)
    a = ap.parse_args(argv)
    for p in (a.view, a.md, a.xlsx):
        if not path_ok(p) or not Path(p).is_file():
            return fail(EXIT_INPUT, f"file not found or path too long: {p}")
    try:
        import openpyxl
    except ImportError:
        return fail(EXIT_DEP, "openpyxl is not installed")
    try:
        view = read_json(a.view)
        c, problems = recount(view)
        md = md_tables(Path(a.md).read_text(encoding="utf-8"))
        wb = openpyxl.load_workbook(a.xlsx, read_only=False, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        return fail(EXIT_INPUT, f"cannot read input: {exc}")
    if "Counts by status and severity" in md:
        check_matrix("md", md["Counts by status and severity"], c, problems)
    else:
        problems.append("md: section 'Counts by status and severity' missing")
    if "Counts by route" in md:
        check_routes("md", md["Counts by route"], c, problems)
    else:
        problems.append("md: section 'Counts by route' missing")
    rule_rows = md.get("Findings by rule", [])
    shown = {r[0]: r[1] for r in rule_rows[1:]}
    for k in set(c["rules"]) | set(shown):
        compare(f"md rule {k}", c["rules"].get(k, 0), shown.get(k, "missing"), problems)
    if "Summary" not in wb.sheetnames or "By rule" not in wb.sheetnames:
        problems.append("xlsx: sheet 'Summary' or 'By rule' missing")
    else:
        ws = wb["Summary"]
        mrows = sheet_block(ws, MATRIX_HEAD)
        if mrows:
            check_matrix("xlsx", mrows, c, problems)
        else:
            problems.append("xlsx: status x severity table not found")
        rrows = sheet_block(ws, "Route")
        if rrows:
            check_routes("xlsx", rrows, c, problems)
        else:
            problems.append("xlsx: route table not found")
        shown = {r[0]: r[1] for r in sheet_block(wb["By rule"], "Rule")[1:]}
        for k in set(c["rules"]) | set(shown):
            compare(f"xlsx rule {k}", c["rules"].get(k, 0), shown.get(k, "missing"), problems)
    if problems:
        emit({"status": "fail", "code": EXIT_SCHEMA, "problems": problems[:30], "count": len(problems)})
        return EXIT_SCHEMA
    emit({"status": "pass", "findings": c["total"], "routes": dict(sorted(c["routes"].items())),
          "rules": len(c["rules"])})
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
