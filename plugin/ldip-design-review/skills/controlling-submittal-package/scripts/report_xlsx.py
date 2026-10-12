#!/usr/bin/env python3
"""Write a readable DRAFT workbook of a skill-6 run for standalone use (owner decision D2, 11 Oct 2026).

Usage: python report_xlsx.py --out <report.xlsx> [--package <package.json>] [--midp-match <midp_match.json>]
           [--results <pkg_results.json>] [--candidates <candidates.json>] [--midp-diff <midp_diff.json>]
           [--title "<text>"]
Sheets (each only when its input is given): Summary · Findings · Register · Stage (planned vs received,
information only) · Not run · MIDP changes. Every sheet is headed "DRAFT — not for issue": candidate
findings are not verified (skill 18) and nothing in the workbook is issued (Human Review Gate).
The workbook is written to .tmp and renamed. It is a view of the JSON files; it adds no data.
Exit: 0 written · 2 input error · 3 openpyxl missing. No network, no shell calls.
"""
__version__ = "0.1.0"

import argparse
import collections
import os
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pkg_common import EXIT_DEP, EXIT_INPUT, EXIT_OK, emit, fail, path_ok, read_json  # noqa: E402

DRAFT = "DRAFT — not for issue. Candidate findings are unverified (run skill 18 before use)."
RIYADH = timezone(timedelta(hours=3))


def sheet(wb, name, headers, rows, widths=None, note=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(name[:31])
    ws.append([DRAFT])
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws.append([note or ""])
    ws.append(headers)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([("; ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v) for v in r])
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths or [18] * len(headers)):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = w
    ws.freeze_panes = "A4"
    if rows:
        ws.auto_filter.ref = f"A3:{chr(64 + len(headers))}{3 + len(rows)}"
    return ws


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--package")
    ap.add_argument("--midp-match")
    ap.add_argument("--results")
    ap.add_argument("--candidates")
    ap.add_argument("--midp-diff")
    ap.add_argument("--title", default="Submittal package control (skill 6)")
    a = ap.parse_args(argv)
    if not a.out.lower().endswith(".xlsx") or not path_ok(a.out):
        return fail(EXIT_INPUT, "--out must be an .xlsx path of at most 250 characters")
    if not any([a.package, a.midp_match, a.results, a.candidates, a.midp_diff]):
        return fail(EXIT_INPUT, "give at least one input file")
    try:
        import openpyxl
        from openpyxl.styles import Font
    except ImportError:
        return fail(EXIT_DEP, "openpyxl is not installed")
    try:
        pkg = read_json(a.package) if a.package else None
        mm = read_json(a.midp_match) if a.midp_match else None
        res = read_json(a.results) if a.results else None
        cand = read_json(a.candidates) if a.candidates else None
        diff = read_json(a.midp_diff) if a.midp_diff else None
    except (OSError, ValueError) as exc:
        return fail(EXIT_INPUT, f"cannot read input: {exc}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([DRAFT])
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws.append([a.title])
    ws["A2"].font = Font(bold=True, size=14)
    info = [("Generated", datetime.now(RIYADH).strftime("%Y-%m-%d %H:%M (+03:00)")),
            ("Skill", f"controlling-submittal-package (report_xlsx {__version__})")]
    if pkg:
        info += [("Site / sub", f"{pkg['site']} / {pkg.get('sub_id')}"), ("Mode", pkg.get("mode")),
                 ("Documents (register rows)", pkg["counts"]["documents"]), ("Files", pkg["counts"]["files"]),
                 ("Files without a document number", pkg["counts"]["files_without_number"]),
                 ("Drawing / transmittal list", (pkg.get("transmittal") or {}).get("path") or "not supplied")]
    if mm:
        for f in mm["midp_files"]:
            info.append(("MIDP", f"{f['number']} — updated {f['updated']} — sha256 {f['sha256'][:16]}…"))
        info.append(("MIDP match levels", ", ".join(f"{k} {v}" for k, v in mm["counts"].items())))
        info.append(("Match key", mm["key_definition"]))
        if mm.get("stage"):
            s = mm["stage"]
            info.append((f"Stage '{s['stage']}' ({s['site']})", f"planned {s['planned']}, received {s['received']} "
                                                                f"— information only"))
    if res:
        info.append(("Rulebook", res.get("rulebook_version")))
        for r in res.get("rules_skipped", []):
            info.append(("Rule skipped", f"{r['rule_id']}: {r['reason']}"))
    if cand is not None:
        info.append(("Candidate findings", len(cand)))
        for k, v in sorted(collections.Counter(f"{c['reference']['rule_id']} · {c['proposed_status']} · {c['severity']}"
                                               for c in cand).items()):
            info.append(("  by rule · status · severity", f"{k}: {v}"))
    ws.append([])
    for k, v in info:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 110
    if cand is not None:
        rows = [[c["finding_id"].split("/")[-1], c["reference"]["rule_id"], c["severity"], c["proposed_status"],
                 c["issue_type"], c["evidence"][0].get("measured_value"), c["location"]["document_no"],
                 (c["location"].get("table") or {}).get("sheet"), c["title"], c["statement"],
                 c["evidence"][0]["extract"][:600]] for c in cand]
        sheet(wb, "Findings", ["ID", "Rule", "Severity", "Proposed status", "Issue type", "Count", "Document no.",
                               "Sheet", "Title", "Statement", "Evidence (samples)"], rows,
              [10, 9, 11, 16, 18, 8, 36, 18, 40, 80, 60])
    if pkg:
        byb = {m["base"]: m for m in (mm or {}).get("matches", [])}
        orig = collections.defaultdict(list)
        for o in pkg.get("dwg_origin", []):
            orig[o["base"]].append(o["origin"])
        rows = []
        for r in pkg["sheet_register"]:
            m = byb.get(r["base"], {})
            mf = (m.get("midp") or {}).get("formats", [])
            rows.append([r["base"], r["numbers"], r["revisions"], r["parse"]["field_count"], r["formats_received"],
                         r["pairing"], m.get("match", "—"), (m.get("midp") or {}).get("number") or m.get("candidates"),
                         (m.get("midp") or {}).get("title"), mf, sorted(set(mf) - set(r["formats_received"])),
                         orig.get(r["base"]), [f["path"] for f in r["files"]]])
        for f in pkg.get("files_without_number", []):
            rows.append(["(no number)", "", "", "", [f.get("format")], "", "", "", "", "", "", "", [f["path"]]])
        sheet(wb, "Register", ["Base number", "Numbers received", "Revisions", "Fields", "Formats received",
                               "DWG/PDF pairing", "MIDP match", "MIDP number / candidates", "MIDP title",
                               "MIDP formats", "Missing vs MIDP", "DWG origin", "Files"], rows,
              [40, 40, 10, 7, 16, 12, 12, 40, 36, 18, 16, 18, 60])
    if mm and mm.get("stage"):
        s = mm["stage"]
        rows = [[p["number"], p["title"], p["sheet"], p["cell"], p["planned"], p["actual"], p["revision"],
                 "yes" if p["received"] else "no"] for p in s["received_list"] + s["planned_not_received"]]
        sheet(wb, "Stage", ["MIDP number", "Title", "MIDP sheet", "Cell", "Planned", "Actual (MIDP)", "Revision (MIDP)",
                            "Received in this package"], rows, [42, 50, 20, 8, 12, 12, 10, 12],
              note=f"Stage '{s['stage']}', site {s['site']}: {s['note']}")
    if res:
        rows = []
        for r in res["results"]:
            if r["result"] == "not_run":
                rows.append([r["rule_id"], r["carrier"], "(whole rule)", r.get("reason")])
            for sc in r.get("subchecks_not_run", []) or []:
                rows.append([r["rule_id"], r["carrier"], sc["subcheck"], sc["reason"]])
        for r in res.get("rules_skipped", []):
            rows.append([r["rule_id"], "—", "(inactive in the rulebook)", r["reason"]])
        sheet(wb, "Not run", ["Rule", "Carrier", "Sub-check", "Reason"], rows, [10, 10, 26, 110])
    if diff:
        rows = []
        for f in diff["files"]:
            rows.append(["FILE", f["new_file"]["name"], "", f"old sha256 {f['old_file']['sha256'][:16]}… → new "
                                                            f"{f['new_file']['sha256'][:16]}…"])
            for k, v in f["identity_changes"].items():
                rows.append(["identity", k, v["old"], v["new"]])
            for k, v in f["stage_changes"].items():
                rows.append(["stages", k, v["old"], v["new"]])
            for k, v in f["rows"].items():
                rows.append(["rows", k, "", v])
            for k, v in f["renumbered_patterns"].items():
                rows.append(["renumbered", k, "", v])
            for k, v in f["format_changes"].items():
                rows.append(["formats", k, "", v])
            for k, v in f["removed_only_by_section"].items():
                rows.append(["removed (by section)", k, "", v])
            for k, v in f["added_only_by_section"].items():
                rows.append(["added (by section)", k, "", v])
            for t in f["title_changes"]:
                rows.append(["title", t["number"], t["old"], t["new"]])
            for x in f["renumbered"]:
                rows.append(["renumbered row", x["title"], x["old"], x["new"]])
        sheet(wb, "MIDP changes", ["Type", "Item", "Old", "New"], rows, [20, 60, 50, 50])
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.name + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)
    emit({"status": "ok", "out": str(out), "sheets": wb.sheetnames})
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
