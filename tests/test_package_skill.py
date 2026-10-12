"""Tests for skill 6 controlling-submittal-package (SPC-0006 §4.3; SPC-0004 §5.1, §5.7, §6; owner
decisions D1-D8, 11 Oct 2026). Synthetic numbers and workbooks only; the real MIDP / rulebook checks run
when LDIP_MIDP_DIR and LDIP_RULEBOOK point at local copies (never in Git)."""
import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[1]
PLUGIN = REPO / "plugin/ldip-design-review"
SK = PLUGIN / "skills/controlling-submittal-package/scripts"
RB_SK = PLUGIN / "skills/applying-ldip-rulebook/scripts"


def run(script, *args, folder=SK):
    p = subprocess.run([sys.executable, str(folder / script), *map(str, args)], capture_output=True, text=True)
    out = p.stdout.strip().splitlines()[-1] if p.stdout.strip() else "{}"
    return p.returncode, json.loads(out)


# ------------------------------------------------------------------ fixtures
def P(v, tag="AGREED", blocking=False):
    return {"value": v, "tag": tag, "blocking": blocking, "source_id": "SRC-T", "clause": "test"}


def rule(rid, title, statement, method="", params=None, code_lists=None, state="active", sev="Minor",
         issue="package_control", default=None):
    return {"rule_id": rid, "title": title, "statement": statement, "method": method, "parameters": params or {},
            "code_lists": code_lists or {}, "state": state, "cap": "Confirmed" if state == "active" else "Probable",
            "severity_default": sev, "issue_type": issue, "owner_role": "SBM", "skill": "controlling-submittal-package",
            "default_status": default or {"script_output": "Confirmed", "dwg_vector": "Confirmed"},
            "skip_subchecks_for": [], "exclusions": [], "inputs": []}


LEGACY = ["PRJ-ENT-ORG-02-OLDX-MPL-SCH-00001"]


def rules_file(tmp_path, drop=()):
    cl = {"program": ["PRJ"], "entity": ["ENT"], "organization": ["ORG"], "package": ["01", "03"],
          "section": ["AA", "QA", "BB", "CC", "GG"], "discipline": ["CIV", "ROD", "GEN"],
          "doc_type": ["DWG", "MOD", "RPT", "LST", "GDL"], "level": ["00", "GF"]}
    rules = [
        rule("PKG-001", "Document number structure", "Drawing numbers 10 fields, documents 9.",
             params={"fields_drawing_model": P(10), "fields_document": P(9), "legacy_numbers": P(LEGACY)}),
        rule("PKG-002", "Field codes from the code lists", "Fields use listed codes.",
             params={"legacy_numbers": P(LEGACY)}, code_lists=cl),
        rule("PKG-003", "Sequence number format and category", "Sequence 5 digits.", method="Regex ^[0-9]{5}$; x",
             params={"category_applies_to_documents": P(None, "GAP")}, code_lists={"seq_category": list("0123456789")}),
        rule("PKG-004", "Revision code format and stage", "Revision per §3.10.", state="advisory",
             method="Regex ^(0[A-Z]|[0-9]{2}|A[0-9])$. Stage not checked.",
             params={"legacy_numbers": P(LEGACY), "stage_consistency": P("series only", "PROPOSED")}),
        rule("PKG-005", "Same number and revision on every carrier", "Carriers agree.",
             params={"revision_in_file_name": P(None, "GAP")}),
        rule("PKG-006", "Section code matches the submittal site", "Section matches site.",
             params={"site_map": P({"QRWN": ["AA", "QA", "GG"], "NRJS": ["BB", "GG"], "NMAR": ["CC"]}),
                     "effective_from": P("2026-09-25"), "legacy_numbers": P(LEGACY)}),
        rule("PKG-007", "Package complete against transmittal and drawing list", "Listed = received.", sev="Major",
             default={"script_output": "Missing information", "xlsx_cell": "Missing information"}),
        rule("PKG-008", "DWG + PDF for every drawing, plus native files", "Formats per MIDP.", sev="Major",
             issue="missing_information", params={"drawing_formats": P(["dwg", "pdf"])},
             default={"script_output": "Missing information"}),
        rule("PKG-009", "No duplicate document number within a package", "No duplicates.", sev="Major"),
        rule("PKG-011", "DWG external references resolved", "Xrefs resolve.", state="advisory", sev="Major",
             default={"dwg_vector": "Not auto-verifiable"}),
    ]
    rules = [r for r in rules if r["rule_id"] not in drop]
    p = tmp_path / "rules_INT.json"
    p.write_text(json.dumps({"agent": "INT", "rulebook_version": "v9", "rules": rules,
                             "rules_skipped": [{"rule_id": "PKG-012", "reason": "inactive: test"}]}))
    return p


STAGES = ["Stage A", "Stage B"]


def midp_book(path, rows, ref="PRJ-ENT-ORG-03-CC-GEN-GDL-00001-001", swap=False, stages=STAGES, updated="2026-10-05"):
    """MIDP-like workbook: identity cells, stage row, label row, header row (B = Document Number)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test MIDP"
    for i, (k, v) in enumerate([("Client", "JV"), ("Author", "Consultant"), ("Date Last Updated", updated),
                                ("Document Reference", ref)], start=4):
        ws.cell(i, 2, k)
        ws.cell(i, 3, v)
    fmts = ["DWG", "PDF", "NWD", "RVT", "NWC", "IFC", "DWFx", "DOCX", "XLSX", "XML", "GDB"]
    if swap:
        fmts[8], fmts[9] = fmts[9], fmts[8]
    hdr = ["", "Document Number", "Document Title", *fmts]
    base = len(hdr) + 1
    for s, name in enumerate(stages):
        c = base + 4 * s
        ws.cell(11, c, "Stage:")
        ws.cell(11, c + 1, name)
        for k, lab in enumerate(["Owner", "Planned Submission Date", "Actual Submission Date", "Revision"]):
            ws.cell(12, c + k, lab)
    for j, h in enumerate(hdr, start=1):
        ws.cell(13, j, h)
    r = 14
    for item in rows:
        if isinstance(item, str):           # section heading
            ws.cell(r, 2, item)
        else:
            num, title, flags, ex = item
            ws.cell(r, 2, num)
            ws.cell(r, 3, title)
            for f in flags:
                ws.cell(r, 4 + fmts.index(f), "X")
            for s, vals in ex.items():
                if s not in stages:
                    continue
                c = base + 4 * stages.index(s)
                for k, v in enumerate(vals):
                    if v:
                        ws.cell(r, c + k, v)
        r += 1
    wb.save(path)
    return path


DW = ["DWG", "PDF", "DWFx"]
MIDP_ROWS = [
    "DRAWINGS",
    ("PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001", "Road plan 1", DW, {"Stage A": ("C", "2026-09-01", "2026-09-08", "00")}),
    ("PRJ-ENT-ORG-03-CC-ROD-DWG-10000-002", "Road plan 2", DW, {"Stage A": ("C", "2026-09-01", "", "")}),
    ("PRJ-ENT-ORG-03-CC-CIV-DWG-20000-001", "Civil section", DW, {"Stage A": ("C", "2026-09-01", "", "")}),
    ("PRJ-ENT-ORG-03-CC-XYZ-DWG-30000-001", "Odd discipline", DW, {"Stage B": ("", "2026-12-01", "", "")}),
    "DOCUMENTS",
    ("PRJ-ENT-ORG-03-CC-GEN-RPT-00001-001", "Design report", ["PDF", "DOCX"], {"Stage A": ("C", "2026-09-01", "", "")}),
]


def make_package(tmp_path):
    pkg = tmp_path / "pkg"
    (pkg / "Drawings").mkdir(parents=True, exist_ok=True)
    (pkg / "Again").mkdir(exist_ok=True)
    files = {
        "Drawings/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001.dwg": "a",          # exact; DWFx missing
        "Drawings/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001.pdf": "b",
        "Again/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001.pdf": "b2",          # PKG-009 different content
        "Drawings/PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000-0A.dwg": "c",      # SRC-02 style -> normalised; PDF+DWFx missing
        "Drawings/PRJ-ENT-ORG-03-CC-ROD-DWG-00-10000-0A - plan.pdf": "d",  # key matches 2 rows -> ambiguous
        "Drawings/PRJ-ENT-ORG-01-AA-CIV-DWG-00-55555-0A.dwg": "e",      # wrong site section; unmatched; PDF missing
        "notes.docx": "f",                                               # no number
        "~$lock.docx": "x",                                              # skipped
    }
    for k, v in files.items():
        (pkg / k).write_text(v)
    lst = tmp_path / "list.csv"
    lst.write_text("Transmittal\nDocument Number,Rev,Title\n"
                   "PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001,,Road plan 1\n"
                   "PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000-0A,0A,Civil\n"
                   "PRJ-ENT-ORG-03-CC-GEN-RPT-00001-00,00,Report\n")
    return pkg, lst


def pipeline(tmp_path, stage="Stage A", with_list=True, extra_register=()):
    pkg, lst = make_package(tmp_path)
    out = tmp_path / "out"
    rules = rules_file(tmp_path)
    book = midp_book(tmp_path / "PRJ-ENT-ORG-03-CC-GEN-LST-00001-001 - MIDP.xlsx", MIDP_ROWS)
    assert run("inventory.py", "--root", pkg, "--site", "NMAR", "--sub", "T3", "--out", out / "files.json")[0] == 0
    args = ["--files", out / "files.json", "--site", "NMAR", "--sub", "T3", "--out", out / "package.json", *extra_register]
    if with_list:
        args += ["--list", lst]
    assert run("build_register.py", *args)[0] == 0
    assert run("read_midp.py", "--file", book, "--out", out / "midp.json")[0] == 0
    margs = ["--package", out / "package.json", "--midp", out / "midp.json", "--rules", rules, "--out", out / "mm.json"]
    if stage:
        margs += ["--stage", stage]
    assert run("match_midp.py", *margs)[0] == 0
    rc, s = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", out / "package.json",
                "--midp-match", out / "mm.json", "--out", out / "res.json")
    assert rc == 0, s
    rc, s = run("build_candidates.py", "--results", out / "res.json", "--rules", rules, "--run", "NMAR-T3-20261011T1600",
                "--site", "NMAR", "--sub", "T3", "--package", out / "package.json", "--out", out / "cand.json")
    assert rc == 0, s
    return out, rules, book


def pats(res_file, rid):
    d = json.loads(Path(res_file).read_text())
    return {p["pattern"]: p for r in d["results"] if r["rule_id"] == rid for p in r.get("patterns", [])}


# ------------------------------------------------------------------ inventory / register
def test_inventory_lists_files_and_skips_locks(tmp_path):
    pkg, _ = make_package(tmp_path)
    rc, s = run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", tmp_path / "inv.json")
    assert rc == 0 and s["files"] == 7 and s["skipped"] == 1 and s["with_number"] == 6
    inv = json.loads((tmp_path / "inv.json").read_text())
    f = next(x for x in inv["files"] if x["name"].endswith("0A - plan.pdf"))
    assert f["number"] == "PRJ-ENT-ORG-03-CC-ROD-DWG-00-10000-0A" and f["format"] == "PDF"


def test_inventory_refuses_output_inside_submittal(tmp_path):
    pkg, _ = make_package(tmp_path)
    assert run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", pkg / "inv.json")[0] == 2
    assert not (pkg / "inv.json").exists()


def test_register_groups_pairs_and_list(tmp_path):
    out, _, _ = pipeline(tmp_path)
    pk = json.loads((out / "package.json").read_text())
    reg = {r["base"]: r for r in pk["sheet_register"]}
    assert reg["PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001"]["parse"]["conforming_count"] is False   # 9 fields for DWG
    civ = reg["PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000"]
    assert civ["revisions"] == ["0A"] and civ["pairing"] == "dwg_only"
    assert [m["number"] for m in pk["missing"]] == ["PRJ-ENT-ORG-03-CC-GEN-RPT-00001-00"]
    assert {u["base"] for u in pk["unlisted"]} == {"PRJ-ENT-ORG-01-AA-CIV-DWG-00-55555", "PRJ-ENT-ORG-03-CC-ROD-DWG-00-10000"}
    assert pk["files_without_number"][0]["path"] == "notes.docx"


# ------------------------------------------------------------------ MIDP
def test_read_midp_stage_marker_vs_stage_name(tmp_path):
    book = midp_book(tmp_path / "m.xlsx", MIDP_ROWS, stages=["Stage 1 Concept", "IFC"])
    run("read_midp.py", "--file", book, "--out", tmp_path / "m.json")
    sh = json.loads((tmp_path / "m.json").read_text())["files"][0]["sheets"][0]
    assert [e["stage"] for e in sh["exchanges"]] == ["Stage 1 Concept", "IFC"]


def test_read_midp_by_labels_with_swapped_columns(tmp_path):
    book = midp_book(tmp_path / "m.xlsx", [("PRJ-ENT-ORG-03-CC-GEN-RPT-00001-001", "R", ["XLSX"], {})], swap=True)
    rc, s = run("read_midp.py", "--file", book, "--out", tmp_path / "m.json")
    assert rc == 0 and s["rows"] == 1
    d = json.loads((tmp_path / "m.json").read_text())
    assert d["rows"][0]["formats"] == ["XLSX"]
    assert d["files"][0]["meta"]["Date Last Updated"]["value"] == "2026-10-05"
    assert [e["stage"] for e in d["files"][0]["sheets"][0]["exchanges"]] == STAGES


def test_read_midp_sections_and_exchanges(tmp_path):
    book = midp_book(tmp_path / "m.xlsx", MIDP_ROWS)
    run("read_midp.py", "--file", book, "--out", tmp_path / "m.json")
    rows = json.loads((tmp_path / "m.json").read_text())["rows"]
    assert len(rows) == 5 and rows[0]["section"] == "DRAWINGS" and rows[-1]["section"] == "DOCUMENTS"
    assert rows[0]["exchanges"][0] == {"index": 1, "stage": "Stage A", "owner": "C", "planned": "2026-09-01",
                                       "actual": "2026-09-08", "revision": "00"}
    assert rows[0]["format_cells"]["DWFx"].endswith(str(rows[0]["row"]))


def test_read_midp_refuses_other_workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "not a MIDP"
    wb.save(tmp_path / "x.xlsx")
    assert run("read_midp.py", "--file", tmp_path / "x.xlsx", "--out", tmp_path / "o.json")[0] == 2


def test_diff_midp_renumbered_titles_formats_stages(tmp_path):
    old = midp_book(tmp_path / "old.xlsx", MIDP_ROWS, updated="2026-09-28", stages=["Old stage", "Stage A", "Stage B"])
    new_rows = [r if isinstance(r, str) else list(r) for r in MIDP_ROWS]
    new_rows[1][0] = "PRJ-ENT-ORG-03-CC-TEL-DWG-10000-001"            # renumbered, same title
    new_rows[3][1] = "Civil cross section"                             # title change
    new_rows[6][2] = ["PDF"]                                           # format change
    new_rows[6][3] = {"Stage A": ("C", "2026-09-01", "2026-09-10", "00")}  # exchange filled
    new = midp_book(tmp_path / "new.xlsx", [r if isinstance(r, str) else tuple(r) for r in new_rows])
    run("read_midp.py", "--file", old, "--out", tmp_path / "o.json")
    run("read_midp.py", "--file", new, "--out", tmp_path / "n.json")
    rc, s = run("diff_midp.py", "--old", tmp_path / "o.json", "--new", tmp_path / "n.json", "--out", tmp_path / "d.json")
    assert rc == 0 and s["renumbered"] == 1 and s["removed_only"] == 0 and s["added_only"] == 0
    f = json.loads((tmp_path / "d.json").read_text())["files"][0]
    assert f["renumbered_patterns"] == {"field 6 ROD->TEL": 1}
    assert f["title_changes"][0]["new"] == "Civil cross section"
    assert f["format_changes"] == {"DOCX+PDF -> PDF": 1}
    assert f["identity_changes"]["Date Last Updated"] == {"old": "2026-09-28", "new": "2026-10-05"}
    assert f["stage_changes"]["Test MIDP"]["new"] == STAGES
    assert f["exchange_changes"]["Stage A | actual | filled"] == 1


def test_match_levels(tmp_path):
    out, _, _ = pipeline(tmp_path)
    mm = {m["base"]: m for m in json.loads((out / "mm.json").read_text())["matches"]}
    assert mm["PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001"]["match"] == "exact"
    assert mm["PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000"]["match"] == "normalised"
    assert mm["PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000"]["midp"]["number"] == "PRJ-ENT-ORG-03-CC-CIV-DWG-20000-001"
    amb = mm["PRJ-ENT-ORG-03-CC-ROD-DWG-00-10000"]
    assert amb["match"] == "ambiguous" and amb["candidates_total"] == 2
    assert mm["PRJ-ENT-ORG-01-AA-CIV-DWG-00-55555"]["match"] == "unmatched"


def test_stage_table_is_information_only(tmp_path):
    out, _, _ = pipeline(tmp_path)
    st = json.loads((out / "mm.json").read_text())["stage"]
    assert st["planned"] == 4 and st["received"] == 2
    assert {p["number"] for p in st["planned_not_received"]} == {"PRJ-ENT-ORG-03-CC-ROD-DWG-10000-002",
                                                                 "PRJ-ENT-ORG-03-CC-GEN-RPT-00001-001"}
    cand = json.loads((out / "cand.json").read_text())
    assert not any("ROD-DWG-10000-002" in c["statement"] for c in cand)


def test_unknown_stage_lists_available(tmp_path):
    pkg, _ = make_package(tmp_path)
    out = tmp_path / "out"
    rules = rules_file(tmp_path)
    book = midp_book(tmp_path / "m.xlsx", MIDP_ROWS)
    run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", out / "f.json")
    run("build_register.py", "--files", out / "f.json", "--site", "NMAR", "--out", out / "p.json")
    run("read_midp.py", "--file", book, "--out", out / "m.json")
    rc, s = run("match_midp.py", "--package", out / "p.json", "--midp", out / "m.json", "--rules", rules,
                "--stage", "Nope", "--out", out / "mm.json")
    assert rc == 2 and s["available_stages"] == STAGES


# ------------------------------------------------------------------ rules on the received files
def test_file_rules(tmp_path):
    out, _, _ = pipeline(tmp_path)
    r = out / "res.json"
    assert pats(r, "PKG-001")["drawing_model_fields_9_expected_10"]["count"] == 1
    assert "no_document_number" in pats(r, "PKG-001")
    p6 = pats(r, "PKG-006")["section_not_for_site"]
    assert p6["count"] == 1 and p6["codes"] == {"AA": 1}
    p7 = pats(r, "PKG-007")
    assert p7["listed_not_received"]["count"] == 1 and p7["received_not_listed"]["count"] == 3
    p8 = pats(r, "PKG-008")
    assert p8["missing_DWFx/midp"]["count"] == 2              # exact + normalised drawings
    assert p8["missing_PDF/midp"]["count"] == 1               # normalised CIV drawing
    assert p8["missing_PDF/rulebook"]["count"] == 1           # unmatched drawing; ambiguous one is PDF-only
    assert pats(r, "PKG-009")["same_number_different_content"]["count"] == 1
    res = json.loads(r.read_text())
    nr = {x["rule_id"]: x for x in res["results"] if x["result"] == "not_run"}
    assert set(nr) == {"PKG-010", "PKG-011"}


def test_candidates_status_severity_and_schema(tmp_path):
    out, rules, _ = pipeline(tmp_path)
    cand = json.loads((out / "cand.json").read_text())
    by = {}
    for c in cand:
        by.setdefault(c["title"].split(":")[-1].strip(), c)
    miss = by["listed not received"]
    assert miss["issue_type"] == "missing_information" and miss["proposed_status"] == "Missing information"
    assert miss["location"]["table"]["cells"] == ["A5"]
    unl = by["received not listed"]
    assert unl["issue_type"] == "package_control" and unl["severity"] == "Minor" and unl["proposed_status"] == "Confirmed"
    m8 = by["missing DWFx/midp"]
    assert m8["evidence"][0]["source_type"] == "xlsx_cell" and m8["evidence"][0]["extract"].endswith("=X")
    assert "MIDP" in m8["reference"]["source_doc"] and "2026-10-05" in m8["reference"]["source_doc"]
    assert run("validate_findings.py", "--findings", out / "cand.json", "--rules", rules)[0] == 0


def test_validate_findings_rejects_cap_breach(tmp_path):
    out, rules, _ = pipeline(tmp_path)
    cand = json.loads((out / "cand.json").read_text())
    cand[0]["reference"]["rule_id"] = "PKG-004"                # advisory -> cap Probable
    cand[0]["proposed_status"] = "Confirmed"
    (out / "bad.json").write_text(json.dumps(cand))
    rc, s = run("validate_findings.py", "--findings", out / "bad.json", "--rules", rules)
    assert rc == 4 and any("cap" in e for e in s["errors"])


def test_without_list_pkg007_not_run(tmp_path):
    out, _, _ = pipeline(tmp_path, with_list=False)
    res = json.loads((out / "res.json").read_text())
    p7 = next(r for r in res["results"] if r["rule_id"] == "PKG-007")
    assert p7["result"] == "not_run" and "K4" in p7["reason"]


def test_title_block_and_xref_from_ingestion(tmp_path):
    pkg, _ = make_package(tmp_path)
    dwg = pkg / "Drawings/PRJ-ENT-ORG-03-CC-CIV-DWG-00-20000-0A.dwg"
    import hashlib
    sha = hashlib.sha256(dwg.read_bytes()).hexdigest()
    ing = tmp_path / "civ.dwg.json"
    ing.write_text(json.dumps({"source": {"sha256": sha, "path": "x"},
                               "title_block": {"document_no": "PRJ-ENT-ORG-03-CC-CIV-DWG-00-29999-0A", "revision": "0B"},
                               "content": {"aecc": {"civil3d_native": True}, "xrefs": [{"name": "BASE", "resolved": False}],
                                           "images": [], "fonts": []}}))
    out, _, _ = pipeline(tmp_path, extra_register=["--ingestion", ing])
    r = out / "res.json"
    p5 = pats(r, "PKG-005")
    assert p5["dwg_title_block_vs_file_name"]["count"] == 1
    assert p5["drawing_list_vs_dwg_title_block"]["count"] == 1
    assert p5["revision_drawing_list_vs_dwg_title_block"]["count"] == 1
    assert pats(r, "PKG-011")["unresolved_xrefs"]["count"] == 1
    pk = json.loads((out / "package.json").read_text())
    assert pk["dwg_origin"][0]["origin"] == "civil3d_native"
    cand = json.loads((out / "cand.json").read_text())
    x = next(c for c in cand if c["reference"]["rule_id"] == "PKG-011")
    assert x["proposed_status"] == "Not auto-verifiable"


def test_legacy_numbers_are_observations(tmp_path):
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    (pkg / f"{LEGACY[0]}.xlsx").write_text("s")
    rules = rules_file(tmp_path)
    run("inventory.py", "--root", pkg, "--site", "QRWN", "--out", tmp_path / "f.json")
    run("build_register.py", "--files", tmp_path / "f.json", "--site", "QRWN", "--out", tmp_path / "p.json")
    run("check_package.py", "--rules", rules, "--site", "QRWN", "--package", tmp_path / "p.json", "--out", tmp_path / "r.json")
    rc, s = run("build_candidates.py", "--results", tmp_path / "r.json", "--rules", rules, "--run", "QRWN-T4-20261011T1600",
                "--site", "QRWN", "--sub", "T4", "--out", tmp_path / "c.json")
    assert rc == 0
    cand = json.loads((tmp_path / "c.json").read_text())
    assert cand and all(c["severity"] == "Observation" for c in cand)
    assert {c["reference"]["rule_id"] for c in cand} == {"PKG-001", "PKG-002", "PKG-004", "PKG-006"}
    assert any("revision missing" in c["title"] for c in cand)


def test_issue_date_before_effective_date_is_legacy(tmp_path):
    out, rules, _ = pipeline(tmp_path)
    run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", out / "package.json",
        "--issue-date", "2026-09-01", "--out", out / "r2.json")
    assert "section_not_for_site/legacy" in pats(out / "r2.json", "PKG-006")


# ------------------------------------------------------------------ MIDP's own numbers (T2 shape)
def test_midp_self_rules_per_sheet(tmp_path):
    rules = rules_file(tmp_path)
    book = midp_book(tmp_path / "PRJ-ENT-ORG-03-CC-GEN-LST-00001-001 - MIDP.xlsx", MIDP_ROWS)
    run("read_midp.py", "--file", book, "--out", tmp_path / "m.json")
    rc, s = run("check_package.py", "--rules", rules, "--site", "NMAR", "--midp", tmp_path / "m.json", "--midp-self",
                "--out", tmp_path / "r.json")
    assert rc == 0
    r = tmp_path / "r.json"
    assert pats(r, "PKG-001")["drawing_model_fields_9_expected_10"]["count"] == 4
    assert pats(r, "PKG-002")["discipline_code_not_listed"]["codes"] == {"XYZ": 1}
    assert pats(r, "PKG-004")["revision_format"]["codes"] == {"001": 1}
    p5 = pats(r, "PKG-005")["file_name_vs_document_reference"]
    assert p5["samples"][0]["extract"] == "C7=PRJ-ENT-ORG-03-CC-GEN-GDL-00001-001"
    res = json.loads(r.read_text())
    assert not any(x["rule_id"] in ("PKG-006", "PKG-007", "PKG-008") for x in res["results"])
    rc, _ = run("build_candidates.py", "--results", r, "--rules", rules, "--run", "NMAR-MIDP-20261011T1600",
                "--site", "NMAR", "--sub", "MIDP", "--out", tmp_path / "c.json")
    assert rc == 0
    cand = json.loads((tmp_path / "c.json").read_text())
    assert all(c["location"]["table"]["sheet"] == "Test MIDP" for c in cand)
    rc, s = run("recheck.py", "--candidates", tmp_path / "c.json", "--root", ".", "--out", tmp_path / "rc.json",
                folder=PLUGIN / "skills/verifying-review-findings/scripts")
    assert rc == 0 and s["results"] == {"agrees": len(cand)}


def test_build_candidates_rejects_bad_run_id(tmp_path):
    out, rules, _ = pipeline(tmp_path)
    rc, _ = run("build_candidates.py", "--results", out / "res.json", "--rules", rules, "--run", "bad",
                "--site", "NMAR", "--sub", "T3", "--out", out / "x.json")
    assert rc == 2 and not (out / "x.json").exists()


def test_report_workbook(tmp_path):
    out, _, book = pipeline(tmp_path)
    run("read_midp.py", "--file", book, "--out", out / "m2.json")
    run("diff_midp.py", "--old", out / "midp.json", "--new", out / "m2.json", "--out", out / "d.json")
    rc, s = run("report_xlsx.py", "--out", out / "r.xlsx", "--package", out / "package.json", "--midp-match", out / "mm.json",
                "--results", out / "res.json", "--candidates", out / "cand.json", "--midp-diff", out / "d.json")
    assert rc == 0 and s["sheets"] == ["Summary", "Findings", "Register", "Stage", "Not run", "MIDP changes"]
    wb = openpyxl.load_workbook(out / "r.xlsx")
    assert all(ws["A1"].value.startswith("DRAFT") for ws in wb.worksheets)
    assert wb["Findings"].max_row == 3 + len(json.loads((out / "cand.json").read_text()))
    assert not list(out.glob("*.tmp*"))


def test_skill_layout_and_frontmatter():
    import yaml
    sk = PLUGIN / "skills/controlling-submittal-package"
    text = (sk / "SKILL.md").read_text(encoding="utf-8")
    fm = yaml.safe_load(text.split("---")[1])
    assert fm["name"] == "controlling-submittal-package" and fm["description"].startswith("LDIP design review (AICC JV):")
    assert len(fm["description"]) <= 1024 and len(text.splitlines()) < 500
    assert len(json.loads((sk / "evals/evals.json").read_text())["evals"]) >= 3
    for s in ("inventory.py", "merge_batches.py", "read_midp.py", "diff_midp.py", "build_register.py",
              "match_midp.py", "check_package.py", "build_candidates.py", "validate_findings.py", "report_xlsx.py"):
        assert f"{s}" in text, s
        src = (sk / "scripts" / s).read_text(encoding="utf-8")
        assert "__version__" in src and "import requests" not in src and "subprocess" not in src


# ------------------------------------------------------------------ real data (local only)
@pytest.mark.skipif(not (os.environ.get("LDIP_MIDP_DIR") and os.environ.get("LDIP_RULEBOOK")),
                    reason="real MIDP / rulebook not set (LDIP_MIDP_DIR, LDIP_RULEBOOK)")
def test_real_midp_t2(tmp_path):
    books = sorted(Path(os.environ["LDIP_MIDP_DIR"]).glob("*.xlsm"))
    run("get_rules.py", "--rulebook", os.environ["LDIP_RULEBOOK"], "--agent", "INT", "--out", tmp_path / "rules.json",
        folder=RB_SK)
    args = [x for b in books for x in ("--file", b)]
    rc, s = run("read_midp.py", *args, "--out", tmp_path / "m.json")
    assert rc == 0 and s["rows"] == 5194
    run("check_package.py", "--rules", tmp_path / "rules.json", "--site", "QRWN", "--midp", tmp_path / "m.json",
        "--midp-self", "--out", tmp_path / "r.json")
    rc, s = run("build_candidates.py", "--results", tmp_path / "r.json", "--rules", tmp_path / "rules.json",
                "--run", "QRWN-MIDP-20261011T1600", "--site", "QRWN", "--sub", "MIDP", "--out", tmp_path / "c.json")
    assert rc == 0 and s["findings"] == 17
    assert s["by_rule"] == {"PKG-001": 3, "PKG-002": 8, "PKG-004": 4, "PKG-005": 2}


def test_merge_batches_run_mode(tmp_path):
    ctx = tmp_path / "0.ai_context/NMAR/T3"
    run_id = "NMAR-T3-20261011T1600"
    for b, name in (("b01", "PRJ-ENT-ORG-03-CC-ROD-DWG-00-10000-0A.dwg"), ("b02", "notes.pdf")):
        (ctx / b).mkdir(parents=True)
        (ctx / b / "batch_manifest.json").write_text(json.dumps({
            "run_id": run_id, "site": "NMAR", "sub_id": "T3", "batch": b, "converter": {"name": "x", "version": "1"},
            "counts": {}, "files": [{"path": f"submittals/NMAR/T3/{name}", "sha256": "a" * 64, "type": "dwg",
                                     "status": "extracted", "output": f"{ctx.as_posix()}/{b}/x.dwg.json", "error": None}]}))
    rc, s = run("merge_batches.py", "--context", ctx, "--run", run_id, "--out", tmp_path / "files.json")
    assert rc == 0 and s["files"] == 2 and s["with_number"] == 1 and s["batches"] == 2
    d = json.loads((tmp_path / "files.json").read_text())
    assert d["mode"] == "run" and d["files"][0]["extraction"]["status"] == "extracted"
    rc, _ = run("merge_batches.py", "--context", ctx, "--run", "NMAR-T3-20261011T1700", "--out", tmp_path / "f2.json")
    assert rc == 2


# ------------------------------------------------------------------ K5 / P-S6-5 / P-S6-6 (11 Oct 2026)
def _ing(path, src, title_block, warnings=()):
    import hashlib
    path.write_text(json.dumps({"source": {"sha256": hashlib.sha256(src.read_bytes()).hexdigest(), "path": src.name},
                                "title_block": title_block, "warnings": [{"code": c, "detail": ""} for c in warnings],
                                "content": {}}))
    return path


def midp_scheme_package(tmp_path):
    pkg = tmp_path / "pkg"
    (pkg / "CAD").mkdir(parents=True)
    (pkg / "PDF").mkdir()
    names = {"CAD/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001-00.dwg": "a", "PDF/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001-00.pdf": "b",
             "CAD/PRJ-ENT-ORG-03-CC-ROD-DWG-00001-001-00.dwg": "c", "PDF/PRJ-ENT-ORG-03-CC-ROD-DWG-00001-001-00.pdf": "d"}
    for k, v in names.items():
        (pkg / k).write_text(v)
    ings = [
        _ing(tmp_path / "a.dwg.json", pkg / "CAD/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001-00.dwg", None,
             ["TITLE_BLOCK_ATTRIBS_UNDECODED"]),
        _ing(tmp_path / "b.pdf.json", pkg / "PDF/PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001-00.pdf",
             {"document_no": "PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001", "revision": "01", "title": "ROAD",
              "date": None, "status": None, "source_ref": "page 1"}),
        _ing(tmp_path / "c.dwg.json", pkg / "CAD/PRJ-ENT-ORG-03-CC-ROD-DWG-00001-001-00.dwg", None, ["TITLE_BLOCK_NOT_FOUND"]),
        _ing(tmp_path / "d.pdf.json", pkg / "PDF/PRJ-ENT-ORG-03-CC-ROD-DWG-00001-001-00.pdf", None, ["TITLE_BLOCK_NOT_FOUND"]),
    ]
    out = tmp_path / "out"
    rules = rules_file(tmp_path)
    assert run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", out / "files.json")[0] == 0
    assert run("build_register.py", "--files", out / "files.json", "--site", "NMAR", "--out", out / "package.json",
               "--ingestion", *ings)[0] == 0
    rc, s = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", out / "package.json",
                "--out", out / "res.json")
    assert rc == 0, s
    return out


def test_midp_scheme_naming_one_pattern(tmp_path):
    out = midp_scheme_package(tmp_path)
    r = out / "res.json"
    assert pats(r, "PKG-001")["midp_scheme_naming"]["count"] == 2
    assert "level_code_not_listed" not in pats(r, "PKG-002")
    assert "sequence_format" not in pats(r, "PKG-003")


def test_title_block_carriers_k5(tmp_path):
    out = midp_scheme_package(tmp_path)
    r = out / "res.json"
    p5 = pats(r, "PKG-005")
    assert "file_name_vs_pdf_title_block" not in p5                       # base = MIDP number
    assert p5["revision_file_name_vs_pdf_title_block"]["count"] == 1      # K5-4: 00 vs 01
    res = json.loads(r.read_text())
    subs = {s["subcheck"]: s["reason"] for x in res["results"] if x["rule_id"] == "PKG-005"
            for s in x.get("subchecks_not_run", [])}
    assert "revision_in_file_name" not in subs
    assert "TITLE_BLOCK_ATTRIBS_UNDECODED" in subs["dwg_title_block_not_read"]   # K5-2
    assert "PRJ-ENT-ORG-03-CC-ROD-DWG-00001-001" in subs["cover_sheets"]          # K5-3
    assert "pdf_title_block_not_read" not in subs


def test_inventory_lists_zip_members(tmp_path):
    import zipfile
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    with zipfile.ZipFile(pkg / "PRJ-ENT-ORG-03-CC-CIV-MOD-00-20000-0A.zip", "w") as z:
        z.writestr("model/PRJ-ENT-ORG-03-CC-CIV-MOD-00-20000-0A.ifc", "ifc")
        z.writestr("inner.zip", "x")
    rc, s = run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", tmp_path / "files.json")
    assert rc == 0 and s["zip_members"] == 2 and s["zip_not_read"] == 1
    inv = json.loads((tmp_path / "files.json").read_text())
    m = next(f for f in inv["files"] if f["kind"] == "zip_member" and f["format"] == "IFC")
    assert m["path"] == "PRJ-ENT-ORG-03-CC-CIV-MOD-00-20000-0A.zip!/model/PRJ-ENT-ORG-03-CC-CIV-MOD-00-20000-0A.ifc"
    assert m["number"] == "PRJ-ENT-ORG-03-CC-CIV-MOD-00-20000-0A" and len(m["sha256"]) == 64
    assert not any(p.name.endswith(".ifc") for p in pkg.rglob("*"))      # nothing extracted


# ------------------------------------------------------------------ 12 Oct 2026: P-S6-7, PKG-007 not received, rules_v1 layouts
def test_unnumbered_files_take_the_zip_or_folder_number(tmp_path):
    import zipfile
    pkg = tmp_path / "pkg"
    (pkg / "Sewer/Model_SWN").mkdir(parents=True)
    with zipfile.ZipFile(pkg / "Sewer/PRJ-ENT-ORG-03-CC-SWN-MOD-00001 Hydraulic Model.zip", "w") as z:
        z.writestr("Model_SWN/model.stsw", "m")
    (pkg / "Sewer/Model_SWN/model.stsw").write_text("m")                    # extracted next to the zip
    (pkg / "Roads/PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001 Roads Model/IFC").mkdir(parents=True)
    (pkg / "Roads/PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001 Roads Model/IFC/readme.txt").write_text("r")
    (pkg / "loose.txt").write_text("x")
    rc, s = run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", tmp_path / "files.json")
    assert rc == 0 and s["number_inherited"] == 3
    inv = {f["path"]: f for f in json.loads((tmp_path / "files.json").read_text())["files"]}
    assert inv["Sewer/Model_SWN/model.stsw"]["number_inherited"]["from"] == "zip"
    assert inv["Roads/PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001 Roads Model/IFC/readme.txt"]["number_inherited"] == {
        "number": "PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001", "from": "folder",
        "source": "Roads/PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001 Roads Model"}
    assert run("build_register.py", "--files", tmp_path / "files.json", "--site", "NMAR", "--list-not-received",
               "--out", tmp_path / "package.json")[0] == 0
    pk = json.loads((tmp_path / "package.json").read_text())
    assert [f["path"] for f in pk["files_without_number"]] == ["loose.txt"]
    assert len(pk["attached_files"]) == 3 and pk["transmittal"] == {"received": False}
    rules = rules_file(tmp_path)
    rc, _ = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", tmp_path / "package.json",
                "--out", tmp_path / "r.json")
    assert rc == 0
    assert pats(tmp_path / "r.json", "PKG-001")["no_document_number"]["count"] == 1
    assert pats(tmp_path / "r.json", "PKG-007")["transmittal_not_received"]["count"] == 1
    rc, s = run("build_candidates.py", "--results", tmp_path / "r.json", "--rules", rules, "--run",
                "NMAR-T3-20261012T0900", "--site", "NMAR", "--sub", "T3", "--package", tmp_path / "package.json",
                "--out", tmp_path / "c.json")
    assert rc == 0, s
    c7 = next(c for c in json.loads((tmp_path / "c.json").read_text()) if c["reference"]["rule_id"] == "PKG-007")
    assert c7["proposed_status"] == "Missing information"


def test_accepted_layouts_midp(tmp_path):
    """rules_v1 PKG-001 accepted_layouts [src02, midp]: MIDP-layout numbers are compliant."""
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    for n in ("PRJ-ENT-ORG-03-CC-ROD-DWG-10000-001-00", "PRJ-ENT-ORG-03-CC-GEN-RPT-00001-001-00",
              "PRJ-ENT-ORG-03-CC-ROD-MOD-90000-001", "PRJ-ENT-ORG-03-CC-ROD-DWG-10000-002-ZZ"):
        (pkg / f"{n}.pdf").write_text(n)
    rules = rules_file(tmp_path)
    d = json.loads(rules.read_text())
    next(r for r in d["rules"] if r["rule_id"] == "PKG-001")["parameters"]["accepted_layouts"] = P(["src02", "midp"])
    rules.write_text(json.dumps(d))
    assert run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", tmp_path / "f.json")[0] == 0
    assert run("build_register.py", "--files", tmp_path / "f.json", "--site", "NMAR", "--out", tmp_path / "p.json")[0] == 0
    pk = json.loads((tmp_path / "p.json").read_text())
    assert "PRJ-ENT-ORG-03-CC-GEN-RPT-00001-001" in [r["base"] for r in pk["sheet_register"]]
    rc, _ = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", tmp_path / "p.json",
                "--out", tmp_path / "r.json")
    assert rc == 0
    r = tmp_path / "r.json"
    assert pats(r, "PKG-001") == {} and "level_code_not_listed" not in pats(r, "PKG-002")
    assert pats(r, "PKG-003") == {}
    p4 = pats(r, "PKG-004")
    assert p4["revision_missing"]["count"] == 1 and p4["revision_format"]["codes"] == {"ZZ": 1}


def test_numbering_scope_midp_deliverables(tmp_path):
    """P-S6-8 (owner 12 Oct 2026, numbering only): PKG-001 structure only for deliverables in the MIDP."""
    pkg = tmp_path / "pkg"
    pkg.mkdir()
    for n in ("PRJ-ENT-ORG-03-CC-GEN-RPT-00001",          # prefix of MIDP GEN-RPT-00001-001 -> checked
              "PRJ-ENT-ORG-03-CC-CIV-RPT-77"):             # not in the MIDP -> not checked
        (pkg / f"{n}.pdf").write_text(n)
    rules = rules_file(tmp_path)
    d = json.loads(rules.read_text())
    next(r for r in d["rules"] if r["rule_id"] == "PKG-001")["parameters"]["numbering_scope"] = P("midp_deliverables")
    rules.write_text(json.dumps(d))
    book = midp_book(tmp_path / "PRJ-ENT-ORG-03-CC-GEN-LST-00001-001 - MIDP.xlsx", MIDP_ROWS)
    assert run("inventory.py", "--root", pkg, "--site", "NMAR", "--out", tmp_path / "f.json")[0] == 0
    assert run("build_register.py", "--files", tmp_path / "f.json", "--site", "NMAR", "--out", tmp_path / "p.json")[0] == 0
    assert run("read_midp.py", "--file", book, "--out", tmp_path / "m.json")[0] == 0
    assert run("match_midp.py", "--package", tmp_path / "p.json", "--midp", tmp_path / "m.json", "--rules", rules,
               "--out", tmp_path / "mm.json")[0] == 0
    mm = {m["base"]: m["in_midp"] for m in json.loads((tmp_path / "mm.json").read_text())["matches"]}
    assert mm["PRJ-ENT-ORG-03-CC-GEN-RPT-00001"]["basis"] == "prefix"
    assert mm["PRJ-ENT-ORG-03-CC-CIV-RPT-77"]["value"] is False
    rc, _ = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", tmp_path / "p.json",
                "--midp-match", tmp_path / "mm.json", "--out", tmp_path / "r.json")
    assert rc == 0
    p1 = pats(tmp_path / "r.json", "PKG-001")
    assert p1["document_fields_8_expected_9"]["count"] == 1
    assert [s["number"] for s in p1["document_fields_8_expected_9"]["samples"]] == ["PRJ-ENT-ORG-03-CC-GEN-RPT-00001"]


def test_pkg013_planned_not_received(tmp_path):
    """PKG-013 (owner activated 12 Oct 2026): MIDP rows planned for the stage but not received."""
    out, base_rules, _ = pipeline(tmp_path)          # stage A: 4 rows planned, 2 not received
    d = json.loads(base_rules.read_text())
    d["rules"].append(rule("PKG-013", "Deliverables planned for the stage received", "Planned = received.", sev="Major",
                           issue="missing_information", default={"script_output": "Missing information",
                                                                 "xlsx_cell": "Missing information"}))
    rules = tmp_path / "rules_13.json"
    rules.write_text(json.dumps(d))
    rc, _ = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", out / "package.json",
                "--midp-match", out / "mm.json", "--out", out / "r13.json")
    assert rc == 0
    p = pats(out / "r13.json", "PKG-013")["planned_not_received"]
    mm = json.loads((out / "mm.json").read_text())["stage"]
    assert p["count"] == len(mm["planned_not_received"]) > 0
    rc, s = run("build_candidates.py", "--results", out / "r13.json", "--rules", rules, "--run", "NMAR-T3-20261012T1100",
                "--site", "NMAR", "--sub", "T3", "--package", out / "package.json", "--out", out / "c13.json")
    assert rc == 0, s
    c = next(x for x in json.loads((out / "c13.json").read_text()) if x["reference"]["rule_id"] == "PKG-013")
    assert c["proposed_status"] == "Missing information" and "(PKG-013)" in c["reference"]["clause"]
    assert c["location"]["table"]["cells"]
    rc, _ = run("check_package.py", "--rules", rules, "--site", "NMAR", "--package", out / "package.json",
                "--out", out / "r13b.json")
    res = json.loads((out / "r13b.json").read_text())
    assert next(x for x in res["results"] if x["rule_id"] == "PKG-013")["result"] == "not_run"
