"""Skill 16 0.1.0 — synthetic T16-1…T16-4 + S32 re-review R1/R3b/R4b/R5/R6b/R9.

SPC-0006 Rev N §4.16. No real project data.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from jsonschema import Draft202012Validator

REPO = Path(__file__).resolve().parents[1]
PLUGIN = REPO / "plugin/ldip-design-review"
SK = PLUGIN / "skills/checking-design-document-consistency/scripts"
REF = PLUGIN / "skills/checking-design-document-consistency/references"
TPL = PLUGIN / "skills/checking-design-document-consistency/templates"
PY = sys.executable


def run(script, *args):
    p = subprocess.run([PY, str(SK / script), *map(str, args)], capture_output=True, text=True)
    line = p.stdout.strip().splitlines()[-1] if p.stdout.strip() else "{}"
    try:
        msg = json.loads(line)
    except json.JSONDecodeError:
        msg = {"raw": p.stdout, "stderr": p.stderr}
    return p.returncode, msg


def write(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=1), encoding="utf-8")
    return path


def load(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def make_boq(path, *, cut_base=5065345, cut_factor=1.05, fill_base=3748575, fill_factor=1.05,
             disposal=800000, imp=120000, ratios=(0.2, 0.8, 0.6),
             ref_alist=5, road_area=105710, kerb=23900,
             nrjs=False):
    """R1 fixture: part sheets 4-EARTHWORKS / 6- Roads + A-List + Dayworks/Electrical traps.

    Cut total lives as helper I13-style base×factor; F13-style common 20 % is a trap portion.
    Dayworks/Electrical contain fill/kerb/import words that must NOT win roles.
    """
    if nrjs:
        cut_base, fill_base, disposal, imp = 37000, 50000, 5000, 36960
        cut_factor = fill_factor = 1.0
        road_area, kerb = 10400, 5278

    wb = openpyxl.Workbook()

    # --- 4-EARTHWORKS (role source) ---
    ws = wb.active
    ws.title = "4-EARTHWORKS"
    ws.append(["Bill No.", "Item No.", "Item Description", "Unit of Measure", "Qty"])
    # Helper total in column I (col 9) row 13 style — put on row 2 col E as total, and a portion
    # Row 2: total excavation with base×factor (the real cut)
    ws.append(["4", "EWK-TOT", "Bulk excavation total", "m3", f"={cut_base}*{cut_factor}"])
    # Row 3: common 20 % portion (trap — must not win cut)
    ws.append(["4", "EWK-COM", "Common excavation 20 %", "m3", f"=E2*{ratios[0]}"])
    ws.append(["4", "EWK-02", "Suitable fill", "m3",
               f"={fill_base}*{fill_factor}" if fill_factor != 1.0 else fill_base])
    ws.append(["4", "EWK-03", "Disposal off site", "m3", disposal])
    ws.append(["4", "EWK-04", "Imported fill", "m3", imp])
    ws.append(["4", "EWK-06", "Rock portion (ratio)", "m3", f"=E2*{ratios[0]}"])
    ws.append(["4", "EWK-07", "Suitable portion (ratio)", "m3", f"=E2*{ratios[1]}"])
    ws.append(["4", "EWK-08", "Disposal portion (ratio)", "m3", f"=E2*{ratios[2]}"])

    # --- 6- Roads and Pavings ---
    ws2 = wb.create_sheet("6- Roads and Pavings")
    ws2.append(["Bill No.", "Item No.", "Item Description", "Unit of Measure", "Qty"])
    ws2.append(["6", "ROD-01", "Road pavement area", "m2",
                f"={road_area}*1.05" if not nrjs else road_area])
    ws2.append(["6", "ROD-02", "Kerb length", "m",
                f"={kerb}*1.05" if not nrjs else kerb])

    # --- A-List (principal) ---
    ws3 = wb.create_sheet("A-List")
    ws3.append(["Bill No.", "Item No.", "Item Description", "Unit of Measure", "Qty"])
    ws3.append(["PQ", "PQ-01", "Principal excavation", "m3", f"={cut_base}*{cut_factor}"])
    for i in range(ref_alist):
        row = 3 + i
        ws3.cell(row=row, column=1, value=f"Part {6 + i}")
        ws3.cell(row=row, column=2, value=f"PQ-ERR{i}")
        ws3.cell(row=row, column=3, value="Principal quantity link")
        ws3.cell(row=row, column=4, value="m3")
        ws3.cell(row=row, column=5, value="=#REF!")

    # --- C- Dayworks (TRAP — must not win import/kerb/road) ---
    ws4 = wb.create_sheet("C- Dayworks")
    ws4.append(["Bill No.", "Item No.", "Item Description", "Unit of Measure", "Qty"])
    ws4.append(["C", "DW-01", "Imported sabkha", "m3", 50])
    ws4.append(["C", "DW-02", "Kerb temporary", "m", 50])
    ws4.append(["C", "DW-03", "Road pavement patch", "m2", 5])

    # --- 11- Electrical Works (TRAP — must not win fill) ---
    ws5 = wb.create_sheet("11- Electrical Works")
    ws5.append(["Bill No.", "Item No.", "Item Description", "Unit of Measure", "Qty"])
    ws5.append(["11", "EL-70", "Earthing fill around electrode", "m3", 1563 if not nrjs else 80])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def make_report_pdf_json(path, pages):
    """pages: list of (page_no, text) or (page_no, spans_list)."""
    page_objs = []
    for entry in pages:
        if len(entry) == 2 and isinstance(entry[1], str):
            page_no, text = entry
            # One span per line to preserve newlines for TOC detection
            spans = []
            y = 20.0
            for line in text.split("\n"):
                x = 10.0
                for w in line.split():
                    spans.append({"text": w, "bbox": [x, y, x + 8, y + 10]})
                    x += 10.0
                y += 14.0
        else:
            page_no, spans = entry
        page_objs.append({"page": page_no, "width_pt": 595, "height_pt": 842,
                          "spans": spans, "ocr_spans": []})
    return write(path, {"content": {"pages": page_objs}})


def schema():
    return json.loads((TPL / "consistency.schema.json").read_text(encoding="utf-8"))


def validate_consistency(doc):
    Draft202012Validator(schema()).validate(doc)
    for f in doc.get("findings") or []:
        assert f.get("reference", {}).get("rule_id") is None
        assert f.get("report_only") is True


# ------------------------------------------------------------------ T16-1 / R1
def test_t16_1_boq_base_factor_and_ratios(tmp_path):
    """T16-1 + R1: cut/fill from part sheets; Dayworks/Electrical traps ignored."""
    xlsx = make_boq(tmp_path / "boq.xlsx")
    out = tmp_path / "boq.json"
    rc, msg = run("read_boq.py", "--xlsx", xlsx, "--out", out,
                  "--document-no", "LDIP-RCRC-SAP-01-QRWN-BOQ-00001")
    assert rc == 0, msg
    boq = load(out)
    cut = boq["summary"]["earthworks"]["cut"]
    fill = boq["summary"]["earthworks"]["fill"]
    imp = boq["summary"]["earthworks"]["import"]
    assert cut["base"] == pytest.approx(5065345)
    assert cut["factor"] == pytest.approx(1.05)
    # Must NOT be the common 20 % portion (≈1,013,069)
    assert cut["base"] > 2_000_000
    assert fill["base"] == pytest.approx(3748575)
    assert imp["base"] == pytest.approx(120000)
    # Trap sheets must not win
    assert imp["base"] != 50
    assert fill["base"] != 1563
    road = boq["summary"]["roads"]["road_area"]
    kerb = boq["summary"]["roads"]["kerb"]
    assert road["base"] == pytest.approx(105710)
    assert kerb["base"] == pytest.approx(23900)
    assert road["base"] != 5 and kerb["base"] != 50
    ratios = set(boq["summary"]["ratios_recorded"])
    assert {0.2, 0.8, 0.6}.issubset(ratios)

    report = write(tmp_path / "report_claims.json", {
        "report_claims_version": "1.0", "sections": [], "claims": [],
        "families_present": [], "pages_read": 1,
    })
    cons = tmp_path / "consistency.json"
    rc, msg = run("check_consistency.py", "--boq", out, "--report", report, "--out", cons)
    assert rc == 0, msg
    doc = load(cons)
    validate_consistency(doc)
    assert any(r["check"] == "C2" and r["result"] == "observation" for r in doc["comparisons"])


def test_r1_nrjs_hard_values(tmp_path):
    """R1: NRJS hard values 37,000 / 50,000 from part sheets."""
    xlsx = make_boq(tmp_path / "boq.xlsx", nrjs=True)
    out = tmp_path / "boq.json"
    rc, msg = run("read_boq.py", "--xlsx", xlsx, "--out", out)
    assert rc == 0, msg
    boq = load(out)
    assert boq["summary"]["earthworks"]["cut"]["base"] == pytest.approx(37000)
    assert boq["summary"]["earthworks"]["fill"]["base"] == pytest.approx(50000)
    assert boq["summary"]["earthworks"]["import"]["base"] == pytest.approx(36960)


# ------------------------------------------------------------------ T16-2 / R3b
def test_t16_2_ref_error_c5_grouped(tmp_path):
    """T16-2 + R3b: handful of grouped C5 findings (A-List), not one per cell."""
    xlsx = make_boq(tmp_path / "boq.xlsx", ref_alist=33)
    out = tmp_path / "boq.json"
    rc, msg = run("read_boq.py", "--xlsx", xlsx, "--out", out)
    assert rc == 0, msg
    boq = load(out)
    assert len(boq["summary"]["cell_errors"]) >= 33
    groups = boq["summary"]["error_groups"]
    alist = next(g for g in groups if g["group"] == "alist")
    assert alist["count"] >= 33

    cons = tmp_path / "consistency.json"
    rc, msg = run("check_consistency.py", "--boq", out, "--out", cons)
    assert rc == 0, msg
    doc = load(cons)
    validate_consistency(doc)
    c5_findings = [f for f in doc["findings"] if (f.get("reference") or {}).get("clause") == "C5"]
    # R3b: handful — far fewer than 33
    assert 1 <= len(c5_findings) <= 5, len(c5_findings)
    assert any("A-List" in f["title"] or "33" in f["title"] or "principal" in f["title"].lower()
               for f in c5_findings)
    # Confirmed for xlsx_cell group
    assert any(f.get("proposed_status") == "Confirmed" for f in c5_findings)


# ------------------------------------------------------------------ T16-3 / R4b / R9
def test_t16_3_balanced_observation(tmp_path):
    """T16-3 + R4b/R9: multi-page TOC skipped; balanced on earthworks body → C3 Obs."""
    xlsx = make_boq(tmp_path / "boq.xlsx", disposal=500000, imp=80000)
    boq_path = tmp_path / "boq.json"
    run("read_boq.py", "--xlsx", xlsx, "--out", boq_path)

    # TOC pages 2–5 with many dot leaders; body earthworks on p10; ICT balanced on p12
    toc_line = "Earthworks ........................ 10\nPavement .......................... 11\nGeotechnical ...................... 9\n"
    toc_pages = [(p, "Contents\n" + toc_line * 3) for p in range(2, 6)]
    pdf = make_report_pdf_json(tmp_path / "report.pdf.json", toc_pages + [
        (10, "9.3.2 Earthworks\nThe cut and fill are balanced across the site grading."),
        (11, "9.10 Pavement\nTotal pavement thickness 550 mm from wearing surface to ToS."),
        (12, "12 ICT\nThe ICT network load is balanced across cabinets."),
    ])
    claims_path = tmp_path / "report_claims.json"
    rc, msg = run("read_report.py", "--text", pdf, "--out", claims_path)
    assert rc == 0, msg
    claims = load(claims_path)
    # R4b: TOC spans multiple pages
    assert len(claims["toc_pages"]) >= 3, claims["toc_pages"]
    # Sections on body pages, not TOC
    sec_pages = {s["section"]: s["page"] for s in claims["sections"]}
    assert sec_pages.get("earthworks") == 10
    assert sec_pages.get("pavement") == 11
    # R9: ICT balanced dropped
    bal = [c for c in claims["claims"] if c["family"] == "balanced"]
    assert len(bal) == 1 and bal[0]["page"] == 10
    assert "balanced" in claims["families_present"]
    # 550 mm found (R5 line rebuild)
    assert any(c.get("value_mm") == pytest.approx(550) for c in claims["claims"]
               if c["family"] == "pavement_depth")

    cons = tmp_path / "consistency.json"
    rc, msg = run("check_consistency.py", "--boq", boq_path, "--report", claims_path, "--out", cons)
    assert rc == 0, msg
    doc = load(cons)
    validate_consistency(doc)
    assert any(r["check"] == "C3" and r["result"] == "observation" for r in doc["comparisons"])


def test_r9_ict_balanced_dropped(tmp_path):
    """R9: ICT 'balanced' without earthworks context is not a claim."""
    pdf = make_report_pdf_json(tmp_path / "r.pdf.json", [
        (10, "9.3 Earthworks Quantities are summarised below."),
        (152, "15 ICT The system is balanced for redundancy."),
    ])
    out = tmp_path / "claims.json"
    rc, msg = run("read_report.py", "--text", pdf, "--out", out)
    assert rc == 0, msg
    claims = load(out)
    assert "balanced" not in claims["families_present"]
    assert not any(c["family"] == "balanced" for c in claims["claims"])


def test_r4b_toc_multipage_sections_on_body(tmp_path):
    """R4b: pages with ≥3 dot-leader lines are TOC; section on first body page."""
    pages = []
    for p in range(2, 10):
        pages.append((p, "\n".join([
            f"Section {i} Earthworks " + "." * 20 + f" {100 + i}" for i in range(5)
        ])))
    pages.append((137, "9.3.2 Earthworks Mass haul and cut fill volumes."))
    pages.append((144, "9.10 Pavement wearing course details."))
    pdf = make_report_pdf_json(tmp_path / "toc.pdf.json", pages)
    out = tmp_path / "claims.json"
    rc, msg = run("read_report.py", "--text", pdf, "--out", out)
    assert rc == 0, msg
    claims = load(out)
    assert all(p in claims["toc_pages"] for p in range(2, 10))
    sec = {s["section"]: s["page"] for s in claims["sections"]}
    assert sec.get("earthworks") == 137
    assert sec.get("pavement") == 144


# ------------------------------------------------------------------ R5 bbox lines + 550 mm
def test_r5_bbox_line_rebuild_550mm(tmp_path):
    """R5: out-of-order spans rebuilt into lines; 550 mm found with pavement context."""
    # Deliberately shuffled: mm before 550, pavement words on different x order
    spans = [
        {"text": "mm", "bbox": [120, 80, 140, 90]},
        {"text": "550", "bbox": [90, 80, 115, 90]},
        {"text": "thickness", "bbox": [50, 60, 100, 70]},
        {"text": "Pavement", "bbox": [10, 40, 60, 50]},
        {"text": "Total", "bbox": [10, 60, 40, 70]},
        {"text": "wearing", "bbox": [10, 80, 50, 90]},
        {"text": "to", "bbox": [150, 80, 165, 90]},
        {"text": "ToS", "bbox": [170, 80, 200, 90]},
        {"text": "design", "bbox": [70, 40, 110, 50]},
    ]
    pdf = make_report_pdf_json(tmp_path / "ooo.pdf.json", [(4, spans)])
    out = tmp_path / "claims.json"
    rc, msg = run("read_report.py", "--text", pdf, "--out", out)
    assert rc == 0, msg
    claims = load(out)
    depths = [c for c in claims["claims"] if c["family"] == "pavement_depth"]
    assert depths and depths[0]["value_mm"] == pytest.approx(550)


# ------------------------------------------------------------------ T16-4
def test_t16_4_pavement_depth_pass_and_flag(tmp_path):
    xlsx = make_boq(tmp_path / "boq.xlsx", disposal=0, imp=0)
    boq_path = tmp_path / "boq.json"
    run("read_boq.py", "--xlsx", xlsx, "--out", boq_path)
    pdf = make_report_pdf_json(tmp_path / "report.pdf.json", [
        (2, "2 Earthworks Site grading overview."),
        (4, "4 Pavement design.\nTotal pavement thickness 550 mm from wearing surface to ToS."),
    ])
    claims_path = tmp_path / "report_claims.json"
    rc, msg = run("read_report.py", "--text", pdf, "--out", claims_path)
    assert rc == 0, msg

    def check_with_median(median, expect_flag):
        profile = write(tmp_path / f"profile_{median}.json", {
            "cut_fill": {"cut_volume_m3": 5.0e6, "fill_volume_m3": 3.5e6},
            "road_top_minus_grading": {
                "n": 100, "median": median, "p5": median - 0.02, "p95": median + 0.02,
                "method": "tin_raster_z",
            },
        })
        cons = tmp_path / f"cons_{median}.json"
        rc, msg = run("check_consistency.py", "--boq", boq_path, "--report", claims_path,
                      "--profile", profile, "--out", cons)
        assert rc == 0, msg
        doc = load(cons)
        validate_consistency(doc)
        c4 = [r for r in doc["comparisons"] if r["check"] == "C4"]
        if expect_flag:
            assert any(r["result"] == "flag" for r in c4)
        else:
            assert any(r["result"] == "pass" for r in c4)

    check_with_median(0.55, expect_flag=False)
    check_with_median(0.40, expect_flag=True)


# ------------------------------------------------------------------ R7 / C6
def test_r7_c1_three_pairs(tmp_path):
    xlsx = make_boq(tmp_path / "boq.xlsx", disposal=0, imp=0)
    boq_path = tmp_path / "boq.json"
    run("read_boq.py", "--xlsx", xlsx, "--out", boq_path)
    profile = write(tmp_path / "profile.json", {
        "cut_fill": {"cut_volume_m3": 4820000.0, "fill_volume_m3": 3230000.0},
    })
    drawing = write(tmp_path / "table.json", {"cut": 7567.55, "fill": 371513.95})
    cons = tmp_path / "consistency.json"
    rc, msg = run("check_consistency.py", "--boq", boq_path, "--profile", profile,
                  "--drawing-table", drawing, "--out", cons)
    assert rc == 0, msg
    labels = {r["quantity"] for r in load(cons)["comparisons"] if r["check"] == "C1"}
    for need in ("cut_boq_vs_model", "cut_boq_vs_drawing", "cut_model_vs_drawing",
                 "fill_boq_vs_model", "fill_boq_vs_drawing", "fill_model_vs_drawing"):
        assert need in labels


def test_c6_information_and_refuses_self(tmp_path):
    xlsx = make_boq(tmp_path / "boq.xlsx")
    boq_path = tmp_path / "boq.json"
    run("read_boq.py", "--xlsx", xlsx, "--out", boq_path)
    reg = write(tmp_path / "register.json", {
        "source": "skill10", "independent": True,
        "roads": {"road_area_m2": 90000.0, "kerb_length_m": 20000.0},
    })
    cons = tmp_path / "c30.json"
    rc, msg = run("check_consistency.py", "--boq", boq_path, "--road-register", reg,
                  "--stage", "30% Concept Design", "--out", cons)
    assert rc == 0, msg
    assert all(r["result"] == "information"
               for r in load(cons)["comparisons"] if r["check"] == "C6")

    self_reg = write(tmp_path / "self.json", {
        "source": "boq",
        "roads": {"road_area_m2": 105710.0, "kerb_length_m": 23900.0},
    })
    cons2 = tmp_path / "cself.json"
    rc, msg = run("check_consistency.py", "--boq", boq_path, "--road-register", self_reg,
                  "--out", cons2)
    assert rc == 0, msg
    assert all(r["result"] == "not_run"
               for r in load(cons2)["comparisons"] if r["check"] == "C6")


# ------------------------------------------------------------------ layout / params
def test_skill_layout_and_version():
    skill = PLUGIN / "skills/checking-design-document-consistency"
    assert 'version: "0.1.0"' in (skill / "SKILL.md").read_text(encoding="utf-8")
    assert (skill / "evals/evals.json").is_file()
    assert 3 <= len(load(skill / "evals/evals.json")["evals"]) <= 5


def test_params_dn5_tag():
    params = load(REF / "params.json")
    tag = params.get("tag") or params["tolerances"].get("tag")
    assert tag and "DN-5" in tag and "PROPOSED" in tag


def test_dn4_no_rule_id(tmp_path):
    xlsx = make_boq(tmp_path / "boq.xlsx", ref_alist=2)
    boq_path = tmp_path / "boq.json"
    run("read_boq.py", "--xlsx", xlsx, "--out", boq_path)
    cons = tmp_path / "c.json"
    rc, msg = run("check_consistency.py", "--boq", boq_path, "--out", cons)
    assert rc == 0, msg
    doc = load(cons)
    assert doc["c7"]["built"] is False
    assert all(f["reference"]["rule_id"] is None for f in doc["findings"])
