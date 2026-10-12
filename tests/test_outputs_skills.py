"""Tests for the standalone versions of skill 22 drafting-review-correspondence and skill 24
generating-review-reports (SPC-0006 §4.8.5, §4.10.4; owner decisions P-S22-1…10, 12 Oct 2026).

Synthetic data only. The two "owner-like" forms are built here from the printed labels of the AICC CRS
and RCRC RFI forms, with invented sample values; the owner's real forms never enter Git. The real-data
check runs only with LDIP_OUTPUT_FINDINGS (and optionally LDIP_CRS_TEMPLATE / LDIP_RFI_TEMPLATE).
"""
import copy
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
PLUGIN = REPO / "plugin/ldip-design-review"
S24 = PLUGIN / "skills/generating-review-reports/scripts"
S22 = PLUGIN / "skills/drafting-review-correspondence/scripts"
PY = sys.executable
RUN = "QRWN-TEST-20261012T1000"
openpyxl = pytest.importorskip("openpyxl")
docx = pytest.importorskip("docx")
DRAFT = "DRAFT — not for issue until approved at the Human Review Gate"


def run(sc, script, *args):
    p = subprocess.run([PY, str(sc / script), *map(str, args)], capture_output=True, text=True)
    lines = p.stdout.strip().splitlines()
    assert lines, p.stderr
    return p.returncode, json.loads(lines[-1])


def finding(n, rule="GIS-001", status="Confirmed", issue="inconsistency", severity="Minor", stage="verified",
            agent="GIS", site="QRWN", run_id=RUN, statement=None, doc="LDIP-TEST-DOC-00001", **loc):
    f = {"schema_version": "1.0", "finding_id": f"{run_id}/{agent}-{n:03d}", "stage": stage, "run_id": run_id,
         "site": site, "sub_id": run_id.split("-")[1], "agent": agent, "title": f"{rule} synthetic finding {n}",
         "statement": statement or f"Synthetic statement {n}: the value differs from the requirement.",
         "issue_type": issue, "severity": severity, "confidence": "High",
         "location": {"document_no": doc, **(loc or {"table": {"sheet": "DLUS", "cells": [f"B{n + 4}"]}})},
         "evidence": [{"evidence_id": "E01", "tier": "deterministic", "source_type": "xlsx_cell",
                       "extract": f"DLUS!B{n + 4}=12{n}"}],
         "reference": {"rule_id": rule, "rulebook_version": "v9", "source_doc": "SRC-TEST", "clause": f"§{n}.1"},
         "proposed_status": status if status in ("Confirmed", "Probable", "Not auto-verifiable",
                                                 "Missing information") else "Probable",
         "owner": {"role": "Lead GIS" if agent == "GIS" else "DMR", "name": None},
         "created_at": "2026-10-12T10:00:00+03:00"}
    if stage != "candidate":
        f["verification"] = {"status": status, "recheck_result": "agrees" if status == "Confirmed" else "not_rechecked",
                             "method": "synthetic re-read", "verified_at": "2026-10-12T10:05:00+03:00"}
        if status == "Not auto-verifiable":
            f["verification"]["manual_check"] = {"what": "check the label", "where": "sheet 1"}
    return f


def write(tmp, name, obj):
    p = tmp / name
    p.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    return p


def view_of(tmp, findings, name="view.json", *extra):
    src = write(tmp, "findings.json", findings)
    rc, out = run(S24, "load_findings.py", "--findings", src, "--out", tmp / name, *extra)
    assert rc == 0, out
    return tmp / name, json.loads((tmp / name).read_text())


# ------------------------------------------------------------------ load_findings + routes (P-S22-2, P-S22-3)
def test_routes_follow_spc0003_table(tmp_path):
    fs = [finding(1),                                                        # Confirmed inconsistency -> CRS
          finding(2, rule="PKG-008", status="Missing information", issue="missing_information", agent="INT"),
          finding(3, issue="er_ambiguity", status="Probable"),               # any er_ambiguity -> RFI
          finding(4, issue="schedule_blocking", status="Confirmed"),
          finding(5, status="Probable"),
          finding(6, status="Not auto-verifiable"),
          finding(7, rule="PKG-001", agent="INT", severity="Observation", issue="package_control",
                  statement="Legacy number has 8 fields; SRC-02 requires 10."),
          finding(8, status="Confirmed", issue="missing_information"),       # no row -> INCONSISTENT
          finding(9, stage="candidate", status="Confirmed")]
    _, v = view_of(tmp_path, fs)
    route = {r["finding_id"].split("-")[-1]: r["route"] for r in v["findings"]}
    assert route == {"001": "CRS", "002": "CRS", "003": "RFI", "004": "RISK_EOT_REGISTER", "005": "ENGINEER_QUEUE",
                     "006": "MANUAL_CHECK", "007": "NONE", "008": "INCONSISTENT", "009": "UNVERIFIED"}
    assert all(r["register_no"] is None for r in v["findings"])
    assert {r["route_kind"] for r in v["findings"] if r["stage"] == "verified"} == {"provisional"}
    assert v["counts"]["stages"] == {"candidate": 1, "verified": 8}


def test_load_refuses_invalid_and_duplicates(tmp_path):
    bad = finding(1)
    bad["verification"]["status"] = "Confirmed"
    bad["confidence"] = "Low"                                   # Confirmed with Low confidence: schema rejects
    rc, out = run(S24, "load_findings.py", "--findings", write(tmp_path, "b.json", [bad]), "--out", tmp_path / "v.json")
    assert rc == 4 and not (tmp_path / "v.json").exists() and out["errors"][0]["finding_id"] == bad["finding_id"]
    rc, out = run(S24, "load_findings.py", "--findings", write(tmp_path, "d.json", [finding(1), finding(1)]),
                  "--out", tmp_path / "v.json")
    assert rc == 2 and "twice" in out["detail"]


def test_legacy_list_from_rulebook(tmp_path):
    f = finding(1, rule="PKG-002", agent="INT", severity="Observation", issue="package_control",
                statement="Number uses discipline codes not in the code list.", doc="OLD-NUM-0001")
    rb = write(tmp_path, "rules.json", {"rulebook_version": "v9", "rules": [
        {"rule_id": "PKG-002", "parameters": {"legacy_numbers": {"value": ["OLD-NUM-0001"], "tag": "AGREED"}}}]})
    _, v = view_of(tmp_path, [f], "v1.json")
    assert v["findings"][0]["route"] == "CRS"                    # without the rulebook: not recognised
    _, v = view_of(tmp_path, [f], "v2.json", "--rulebook", rb)
    assert v["findings"][0]["route"] == "NONE"


# ------------------------------------------------------------------ skill 24 report (T1-T3)
def report(tmp, view, *extra):
    rc, out = run(S24, "build_report.py", "--view", view, "--out-md", tmp / "report.md",
                  "--out-xlsx", tmp / "report.xlsx", *extra)
    assert rc == 0, out
    return out


def test_t1_zero_findings_lists_rules_run_and_skipped(tmp_path):
    view, _ = view_of(tmp_path, [])
    res = write(tmp_path, "results.json", {"results": [
        {"rule_id": "GIS-001", "result": "pass"}, {"rule_id": "GIS-002", "result": "pass"},
        {"rule_id": "GIS-009", "result": "not_run", "reason": "plot layers [GAP]"}],
        "rules_skipped": [{"rule_id": "GIS-004", "reason": "inactive: layer mapping GAP"}]})
    log = tmp_path / "run.jsonl"
    log.write_text(json.dumps({"ts": "t", "agent": "intake-batch", "instance": "b01", "event": "fail",
                               "detail": {"step": "1a"}}) + "\n" +
                   json.dumps({"ts": "t", "agent": "gis-cadastral", "event": "rule_skipped",
                               "detail": {"rule_id": "GIS-017", "reason": "inactive"}}) + "\n")
    out = report(tmp_path, view, "--results", res, "--run-log", log)
    assert out["findings"] == 0 and out["rules_run"] == 2 and out["not_run"] == 3 and out["failed_steps"] == 1
    md = (tmp_path / "report.md").read_text()
    assert "**No findings.**" in md and DRAFT in md and "GIS-001, GIS-002" in md
    for rid in ("GIS-004", "GIS-009", "GIS-017"):
        assert rid in md
    wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
    assert all(ws["A1"].value.startswith(DRAFT) for ws in wb.worksheets)
    assert run(S24, "check_report.py", "--view", view, "--md", tmp_path / "report.md",
               "--xlsx", tmp_path / "report.xlsx")[0] == 0


def qrwn_like():
    """SPC-0004 §8.1 QRWN expectation shape: 8 Confirmed findings, Major 1, Minor 7."""
    rules = ["GIS-001", "GIS-003", "GIS-012", "GIS-013", "GIS-013", "GIS-014", "GIS-015", "GIS-015"]
    return [finding(i, rule=r, severity="Major" if r == "GIS-001" else "Minor") for i, r in enumerate(rules, 1)]


def test_t2_counts_match(tmp_path):
    view, _ = view_of(tmp_path, qrwn_like())
    report(tmp_path, view)
    rc, out = run(S24, "check_report.py", "--view", view, "--md", tmp_path / "report.md",
                  "--xlsx", tmp_path / "report.xlsx")
    assert rc == 0 and out["findings"] == 8 and out["routes"] == {"CRS": 8}
    md = (tmp_path / "report.md").read_text()
    assert "| Confirmed | 0 | 1 | 7 | 0 | 8 |" in md and "| GIS-013 | 2 |" in md


def test_t3_hand_edited_counts_fail(tmp_path):
    view, _ = view_of(tmp_path, qrwn_like())
    report(tmp_path, view)
    md = tmp_path / "report.md"
    good = md.read_text()
    md.write_text(good.replace("| Confirmed | 0 | 1 | 7 | 0 | 8 |", "| Confirmed | 0 | 1 | 6 | 0 | 7 |"))
    rc, out = run(S24, "check_report.py", "--view", view, "--md", md, "--xlsx", tmp_path / "report.xlsx")
    assert rc == 4 and any("Confirmed / Minor" in p for p in out["problems"])
    md.write_text(good)
    wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
    ws = wb["By rule"]
    for row in ws.iter_rows():
        if row[0].value == "GIS-013":
            row[1].value = 3
    wb.save(tmp_path / "report.xlsx")
    rc, out = run(S24, "check_report.py", "--view", view, "--md", md, "--xlsx", tmp_path / "report.xlsx")
    assert rc == 4 and any("xlsx rule GIS-013" in p for p in out["problems"])
    report(tmp_path, view)
    src = tmp_path / "findings.json"
    src.write_text(src.read_text() + " ")                      # input changed after the view was built
    rc, out = run(S24, "check_report.py", "--view", view, "--md", md, "--xlsx", tmp_path / "report.xlsx")
    assert rc == 4 and any("changed" in p for p in out["problems"])


# ------------------------------------------------------------------ skill 22 drafts (T1-T3, P-S22-4/5)
def drafts(tmp, view, *extra):
    rc, out = run(S22, "draft_crs.py", "--view", view, "--out", tmp / "drafts.json", *extra)
    return rc, out, (json.loads((tmp / "drafts.json").read_text()) if rc == 0 else None)


def test_t1_three_confirmed_crs_rows(tmp_path):
    fs = [finding(1, rule="GIS-001", severity="Major"), finding(2, rule="GIS-012"),
          finding(3, rule="GIS-014", plot_ids=["1958", "1959"])]
    view, _ = view_of(tmp_path, fs)
    rc, out, d = drafts(tmp_path, view)
    assert rc == 0 and out["crs_rows"] == 3 and out["crs_sheets"] == 1
    rows = d["crs"][0]["rows"]
    for r, f in zip(rows, fs):
        assert r["comment"].startswith(f["statement"])
        assert f["reference"]["rule_id"] in r["comment"] and f["reference"]["clause"] in r["comment"]
        assert r["location"].startswith("LDIP-TEST-DOC-00001") and r["comment_status"] == "Open"
    assert "plots 1958, 1959" in rows[2]["location"]
    assert d["marking"].startswith(DRAFT) and d["crs"][0]["header"]["document_no"] == "LDIP-TEST-DOC-00001"
    assert run(S22, "validate_drafts.py", "--drafts", tmp_path / "drafts.json", "--view", view)[0] == 0


def test_t2_er_ambiguity_gives_one_rfi(tmp_path):
    f = finding(1, rule=None, issue="er_ambiguity", status="Probable", agent="INT",
                statement="ER clause numbers 7.11 to 7.15 are used twice; the governing clause cannot be identified.",
                doc="ER-LDIP", package_level=True)
    f["reference"] = {"rule_id": None, "source_doc": "Employer's Requirements — LDIP", "clause": "§7.11–7.15"}
    view, _ = view_of(tmp_path, [f])
    rc, out, d = drafts(tmp_path, view)
    assert rc == 0 and out["rfi"] == 1 and out["crs_rows"] == 0
    r = d["rfi"][0]
    assert r["specification"] == "§7.11–7.15" and "used twice" in r["information_requested"]
    assert r["location"] == "AL-QAIRAWAN" and r["rfi_ref"].startswith("DRAFT")


def test_t3_probable_not_drafted_but_queued(tmp_path):
    view, _ = view_of(tmp_path, [finding(1, status="Probable"), finding(2, status="Not auto-verifiable")])
    rc, out, d = drafts(tmp_path, view)
    assert rc == 0 and out["crs_rows"] == 0 and out["rfi"] == 0
    assert [e["finding_id"] for e in d["engineer_queue"]] == [f"{RUN}/GIS-001"]
    assert d["manual_check"][0]["manual_check"]["what"] == "check the label"


def test_candidates_refused_or_left_out(tmp_path):
    view, _ = view_of(tmp_path, [finding(1, stage="candidate"), finding(2, stage="candidate")], "v1.json")
    rc, out, _ = drafts(tmp_path, view)
    assert rc == 5 and "skill 18" in out["detail"]
    view, _ = view_of(tmp_path, [finding(1, stage="candidate"), finding(2)], "v2.json")
    rc, out, d = drafts(tmp_path, view)
    assert rc == 0 and out["crs_rows"] == 1
    assert d["not_drafted"] == [{"finding_id": f"{RUN}/GIS-001", "route": "UNVERIFIED",
                                 "reason": "candidate — run skill 18 (verifying-review-findings) first"}]


def test_validate_catches_edits(tmp_path):
    view, _ = view_of(tmp_path, qrwn_like())
    rc, _, d = drafts(tmp_path, view)
    p = tmp_path / "drafts.json"
    edited = copy.deepcopy(d)
    edited["crs"][0]["rows"][0]["comment"] = "Reworded by hand: please fix the plot numbers."
    p.write_text(json.dumps(edited))
    rc, out = run(S22, "validate_drafts.py", "--drafts", p, "--view", view)
    assert rc == 4 and any("crs/0/rows/0/comment" in x for x in out["problems"])
    edited = copy.deepcopy(d)
    edited["crs"][0]["rows"].pop()
    edited["counts"]["crs_rows"] -= 1
    p.write_text(json.dumps(edited))
    rc, out = run(S22, "validate_drafts.py", "--drafts", p, "--view", view)
    assert rc == 4 and any("7 item(s) in the drafts, 8 expected" in x for x in out["problems"])
    p.write_text(json.dumps(d))
    view.write_text(view.read_text() + " ")
    rc, out = run(S22, "validate_drafts.py", "--drafts", p, "--view", view)
    assert rc == 4 and any("SHA-256" in x for x in out["problems"])


def test_header_override_and_bad_key(tmp_path):
    view, _ = view_of(tmp_path, [finding(1, site="NMAR", run_id="NMAR-30pct-20261012T1000")])
    rc, _, d = drafts(tmp_path, view, "--header", "reviewer_org=AICC JV")
    h = d["crs"][0]["header"]
    assert rc == 0 and h["reviewer_org"] == "AICC JV" and h["section"] == "Package-3"
    assert run(S22, "validate_drafts.py", "--drafts", tmp_path / "drafts.json", "--view", view)[0] == 0
    rc, out, _ = drafts(tmp_path, view, "--header", "colour=red")
    assert rc == 2


# ------------------------------------------------------------------ render_templates (owner-like and neutral forms)
def owner_like_crs(path):
    """A form with the printed labels of the AICC CRS and invented sample values (synthetic)."""
    from openpyxl.styles import Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    thin = Side(style="thin")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RCS"
    ws["A5"] = "Comments Review Sheet (CRS)"
    left = [("CRS Ref:", "SAMPLE-CRS-001"), ("Program", "LDIP"), ("Project:", "Sample project"),
            ("Section", "Package-9"), ("Sub-Section", "XX"), ("PMCM", "Sample PMCM"), ("Transmittal No.:", "T-1"),
            ("Subject:", "Sample subject"), ("Document No:", " SAMPLE-DOC "), ("Reviewer/s Organization:", "Org"),
            ("Reviewer/s Name:", "Sample Name"), ("Contractor:", "Sample contractor")]
    for i, (lab, val) in enumerate(left):
        ws.cell(6 + i, 2, lab)
        ws.cell(6 + i, 3, val)
    for i, lab in enumerate(["This submission is returned (Review Code):", "Date Comments Issued to Contractor: ",
                             "Date of receiving Contractor's Response: "] * 2):
        ws.cell(7 + i, 6, lab)
        ws.cell(7 + i, 7, "C-Revise and Resubmit" if i == 0 else "2026-08-26")
        ws.merge_cells(start_row=7 + i, start_column=7, end_row=7 + i, end_column=8)
    for j, h in enumerate(["#", "Page/Para/Clause/ Doc. No. Etc.", "Review Comment", "Reviewer\nInitial",
                           "Review No. ", "Contractor's Response", "Responder\nInitial", "Status Code",
                           "Comment Status (Closed/Open)"], 1):
        ws.cell(20, j, h)
    for r in range(21, 37):
        for j in range(1, 10):
            ws.cell(r, j).border = Border(thin, thin, thin, thin)
    for n in range(1, 4):
        ws.cell(21 + n, 1, n)
        ws.cell(21 + n, 3, f"Old sample comment {n}")
        ws.cell(21 + n, 4, "XX")
        ws.cell(21 + n, 8, "C")
    for i, (lab, name) in enumerate([("Reviewer 1: PMCM", "Sample Reviewer"), ("Reviewer 2: PMCM", "Second Name")]):
        ws.cell(38 + i, 1, lab)
        ws.merge_cells(start_row=38 + i, start_column=1, end_row=38 + i, end_column=2)
        ws.cell(38 + i, 3, name)
        ws.cell(38 + i, 4, "SR")
        ws.cell(38 + i, 5, "Sign:  ")
        ws.cell(38 + i, 6, "SR")
        ws.cell(38 + i, 7, "Date:  26/08/2026")
    ws.cell(43, 1, "Reviewed Action: A-No Objection B-No Objection with Comments C-Revise and Resubmit D-Rejected")
    dv = DataValidation(type="list", formula1='"Open,Closed,Info"')
    dv.add("I21:I36")
    ws.add_data_validation(dv)
    ws.print_area = "A1:I43"
    wb.save(path)
    return path


def owner_like_rfi(path):
    from docx.shared import Pt
    d = docx.Document()
    t = d.add_table(rows=5, cols=4)
    t.cell(0, 0).merge(t.cell(0, 3)).text = "REQUEST FOR INFORMATION (RFI)"
    vals = [("RFI Reference Number:", "SAMPLE-RFI-00004", "Rev. No.", "00"),
            ("Discipline:", "Roads", "Issue Date:", "23-Jul-2026"), ("Enclosure/s:", "No", "Reply Required Date:", ""),
            ("Priority:", "High", "Reply Required Date:", "")]
    for r, row in enumerate(vals, 1):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    t.cell(1, 1).paragraphs[0].runs[0].font.size = Pt(11)
    d.add_paragraph("SUBJECT:")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Sample subject"
    d.add_paragraph("RFI REFERENCES:")
    tb = d.add_table(rows=4, cols=2)
    for i, (a, b) in enumerate([("Drawing No(s).:", ""), ("Document:", ""), ("Specification:", ""),
                                ("Location:", "SAMPLE SITE")]):
        tb.cell(i, 0).text, tb.cell(i, 1).text = a, b
    d.add_paragraph("REASON FOR RFI:")
    tb = d.add_table(rows=2, cols=3)
    tb.cell(0, 0).text, tb.cell(0, 1).text, tb.cell(0, 2).text = ("☒ Insufficient Information",
                                                                   "☐ Engineering Conflict", "☐ Alternate Proposal")
    tb.cell(1, 0).merge(tb.cell(1, 2)).text = "☐ Other, provide explanation:"
    d.add_paragraph("INFORMATION REQUESTED:")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Sample old question about a boundary."
    d.add_paragraph("SUBMITTED BY:")
    tb = d.add_table(rows=2, cols=2)
    tb.cell(0, 0).text, tb.cell(0, 1).text = "Contractor:", "JV LADUN & AL-AYUNI"
    tb.cell(1, 0).text, tb.cell(1, 1).text = "Signature", "Date: 23-Jul-2026"
    d.add_paragraph("RESPONSE:")
    tb = d.add_table(rows=3, cols=2)
    tb.cell(0, 0).text, tb.cell(0, 1).text = "Designer Response:", "PMCM Response:"
    tb.cell(1, 0).text, tb.cell(1, 1).text = "Old designer answer", "Old PMCM answer"
    tb.cell(2, 0).text, tb.cell(2, 1).text = "Designer Stamp/Signature (if applicable):", "PMCM Stamp/Signature:"
    tb = d.add_table(rows=3, cols=2)
    tb.cell(0, 0).merge(tb.cell(0, 1)).text = "ATTACHED DOCUMENTS:"
    tb.cell(1, 0).text = "old-attachment.pdf"
    d.save(path)
    return path


def all_text_docx(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs] + [s.header.paragraphs[0].text for s in d.sections]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def all_values_xlsx(path):
    wb = openpyxl.load_workbook(path)
    return wb, "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value)


@pytest.mark.parametrize("owner_form", [True, False])
def test_render_crs_and_rfi(tmp_path, owner_form):
    fs = [finding(i, rule="GIS-013", statement=f"Statement {i}: GFA differs from area x FAR.") for i in range(1, 19)]
    rfi = finding(30, rule=None, issue="er_ambiguity", status="Probable", agent="INT", doc="ER-LDIP",
                  statement="ER clause numbers 7.11 to 7.15 are used twice in the ER.", package_level=True)
    rfi["reference"] = {"rule_id": None, "source_doc": "Employer's Requirements — LDIP", "clause": "§7.11–7.15"}
    view, _ = view_of(tmp_path, fs + [rfi])
    rc, _, d = drafts(tmp_path, view)
    assert rc == 0
    extra, before = [], {}
    if owner_form:
        extra = ["--crs-template", owner_like_crs(tmp_path / "forms_crs.xlsx"),
                 "--rfi-template", owner_like_rfi(tmp_path / "forms_rfi.docx")]
        before = {p: Path(p).read_bytes() for p in extra[1::2]}
    out_dir = tmp_path / "out"
    rc, out = run(S24, "render_templates.py", "--drafts", tmp_path / "drafts.json", "--out-dir", out_dir, *extra)
    assert rc == 0, out
    assert out["crs_files"] == 1 and out["rfi_files"] == 1
    wb, text = all_values_xlsx(out_dir / "CRS-QRWN-TEST-DRAFT.xlsx")
    ws = wb.worksheets[0]
    assert DRAFT in text and ws.oddHeader.center.text == DRAFT
    for old in ("Old sample comment", "Sample Reviewer", "Second Name", "SAMPLE-CRS-001", "Sample subject",
                "Sample PMCM", "Sample Name", "2026-08-26"):
        assert old not in text, old
    codes = [ws.cell(c.row, 7).value for c in ws["F"] if "review code" in str(c.value or "").lower()]
    assert codes and set(codes) == {"Select Status Code"}      # the reviewers choose the review code
    numbers = [c.value for c in ws["A"] if isinstance(c.value, int)]
    assert numbers == list(range(1, 19))                         # 18 rows: 15 in the form + 3 inserted
    assert "Statement 18: GFA differs" in text and "Draft trace" in wb.sheetnames
    footer = next(c.row for c in ws["A"] if str(c.value or "").startswith("Reviewer 1"))
    assert footer > max(c.row for c in ws["A"] if isinstance(c.value, int))
    if owner_form:
        assert ws.print_area.endswith("$46")                    # print area moved with the inserted rows
        assert any(str(r.max_row) == "39" for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges)
    rtext = all_text_docx(out_dir / "RFI-QRWN-01-DRAFT.docx")
    assert DRAFT in rtext and "used twice in the ER" in rtext and "§7.11–7.15" in rtext
    for old in ("SAMPLE-RFI-00004", "Sample subject", "Sample old question", "23-Jul-2026", "☒",
                "Old designer answer", "old-attachment.pdf", "SAMPLE SITE"):
        assert old not in rtext, old
    assert all(Path(p).read_bytes() == b for p, b in before.items())   # the owner's forms are never changed


def test_render_refuses_damaged_template_and_same_folder(tmp_path):
    view, _ = view_of(tmp_path, [finding(1)])
    drafts(tmp_path, view)
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a workbook")
    rc, out = run(S24, "render_templates.py", "--drafts", tmp_path / "drafts.json", "--out-dir", tmp_path / "o",
                  "--crs-template", bad)
    assert rc == 2 and "unreadable" in out["detail"]
    tpl = owner_like_crs(tmp_path / "t.xlsx")
    rc, out = run(S24, "render_templates.py", "--drafts", tmp_path / "drafts.json", "--out-dir", tmp_path,
                  "--crs-template", tpl)
    assert rc == 2 and "another folder" in out["detail"]


# ------------------------------------------------------------------ real data (never in Git)
@pytest.mark.skipif(not os.environ.get("LDIP_OUTPUT_FINDINGS"), reason="LDIP_OUTPUT_FINDINGS not set (real data)")
def test_real_chain(tmp_path):
    files = os.environ["LDIP_OUTPUT_FINDINGS"].split(os.pathsep)
    rc, out = run(S24, "load_findings.py", "--findings", *files, "--out", tmp_path / "view.json")
    assert rc == 0, out
    report(tmp_path, tmp_path / "view.json")
    assert run(S24, "check_report.py", "--view", tmp_path / "view.json", "--md", tmp_path / "report.md",
               "--xlsx", tmp_path / "report.xlsx")[0] == 0
    rc, out, _ = drafts(tmp_path, tmp_path / "view.json")
    assert rc in (0, 5), out
    if rc == 0:
        assert run(S22, "validate_drafts.py", "--drafts", tmp_path / "drafts.json", "--view",
                   tmp_path / "view.json")[0] == 0
        extra = []
        for env, flag in (("LDIP_CRS_TEMPLATE", "--crs-template"), ("LDIP_RFI_TEMPLATE", "--rfi-template")):
            if os.environ.get(env):
                extra += [flag, os.environ[env]]
        rc, out = run(S24, "render_templates.py", "--drafts", tmp_path / "drafts.json", "--out-dir", tmp_path / "o",
                      *extra)
        assert rc == 0, out
